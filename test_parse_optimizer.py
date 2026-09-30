"""Self-check: python test_parse_optimizer.py"""
import datetime as dt
import json
import tempfile
from pathlib import Path

import openpyxl

import parse_optimizer as p

U, CLIENT = 10_000_000, "ClienteSecreto"


def smart_planner(path, n, tie=False, skew_tp=None, extra_col=False):
    """Libro Smart Planner sintético y consistente con n touchpoints."""
    names = [f"TP{i}" for i in range(n)]
    inv = [1_000_000 * (i + 1) for i in range(n)]
    if tie:
        inv[1] = inv[0]  # misma inversión: el orden de escenarios debe resolverse por GRPs
    cpgrp = [10_000 + 1_000 * i for i in range(n)]
    grps = [a / c for a, c in zip(inv, cpgrp)]
    imps = [g * U / 1e5 for g in grps]
    reach = [0.05 + 0.02 * i for i in range(n)]
    B, G, I = sum(inv), sum(grps), sum(imps)
    net = max(reach) + 0.01

    wb = openpyxl.Workbook()
    ex = wb.active
    ex.title = "Exercise"
    ex["B2"] = "SCOPE"
    ex.merge_cells("B2:C2")
    for r, (k, v) in enumerate([("Country:", "MEXICO"), ("Advertiser:", CLIENT), ("Campaign period:", "2026-10-01, 2026-10-31"),
                                ("Budget:", str(B)), ("Currency:", "MXN"), ("Target:", "P18-35"), ("Target description:", ""),
                                ("Age:", "18, 35"), ("Target size:", str(U)), ("Date:", dt.datetime(2026, 9, 21))], start=3):
        ex.cell(r, 2, k)
        ex.cell(r, 3, v)

    mmt = wb.create_sheet("Media Mix table")
    mmt.append(["Touchpoint", "Budget", "% Budget", "GRPs", "Impressions (000)", "% GRPs / Imps", "Classical Reach",
                "OTS/Freq", "Cost Per Reach Point", "CPGRP", "CPM"])
    mmt.append(["Total", B, 1, G, I, 1, net, G / 100 / net, B / (net * 100), B / G, B / I])
    for i, nm in enumerate(names):
        a = inv[i] * (1.1 if nm == skew_tp else 1)
        mmt.append([nm, a, inv[i] / B, grps[i], imps[i], imps[i] / I, reach[i], grps[i] / 100 / reach[i],
                    inv[i] / (reach[i] * 100), inv[i] / grps[i], inv[i] / imps[i]])

    rf = wb.create_sheet("Reach & Frequency")
    rf.append(["Touchpoint"] + [f"Freq: {k}" for k in range(1, 11)])
    for nm, r in [("Total", net)] + list(zip(names, reach)):
        rf.append([nm] + [r * 0.9 ** k for k in range(10)])
    for title, header, rows in [
        ("Costs", ["Touchpoint", "CPGRP", "CPM"], [[nm, cpgrp[i], inv[i] / imps[i]] for i, nm in enumerate(names)]),
        ("Restrictions", ["Touchpoint", "MIN", "MAX"], [[nm, 0, 100] for nm in names]),
        ("Curve Parameters", ["Touchpoint", "Max. Reach", "Inflection", "Speed"], [[nm, 10, 1, 5] for nm in names]),
        ("TP Colors", ["Touchpoint", "Color"], [[nm, f"#00000{i}"] for i, nm in enumerate(names)]),
        ("Cross Indices", ["Touchpoint 1", "Touchpoint 2", "Index"],
         [[a, b, 100 + i] for i, (a, b) in enumerate((a, b) for j, a in enumerate(names) for b in names[j + 1:])]),
        ("Curves Estimated", ["Touchpoint", "GRPs", "REACH "], [[nm, g, g / 10] for nm in names for g in (0, 10, 20)]),
        ("Datapoints", ["Touchpoint", "GRPs", "REACH "], [[nm, g, g / 10] for nm in names for g in (0, 5)]),
        ("Plan Cost table", ["Cost per", "MXN"], [["Reach Point", B / (net * 100)], ["Total CPGRP", B / G], ["Total CPM", B / I]]),
    ]:
        ws = wb.create_sheet(title)
        for row in [header] + rows:
            ws.append(row)

    bs = wb.create_sheet("Budget Scenarios")
    order = list(reversed(range(n)))  # orden distinto al de Media Mix table
    hdr = []
    for pre in ("budget", "grps", "imps", "share", "reach"):
        hdr += [f"{pre}_Total"] + [f"{pre}_{names[j]}" for j in order]
    bs.append(hdr + (["extra"] if extra_col else []))
    for f in [x / 10 for x in range(5, 16)]:
        row = []
        for tot, per in [(B * f, [inv[j] * f for j in order]), (G * f, [grps[j] * f for j in order]),
                         (I * f, [imps[j] * f for j in order]), (1, [imps[j] / I for j in order]),
                         (net, [reach[j] for j in order])]:
            row += [tot] + per
        bs.append(row + ([0] if extra_col else []))
    wb.save(path)
    return names, inv, grps


def spectrum(path, title, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in [("Application:", "Havas Spectrum"), ("Date:", "25/09/2026"), ("Title:", title), (None, None),
              (None, "Header")] + rows:
        ws.append(r)
    wb.save(path)


def load(out, key):
    return json.loads((out / f"{key}.canonical.json").read_text())


def levels(m):
    return {v["regla"]: v["nivel"] for v in m["validaciones"]}


def test():
    tmp = Path(tempfile.mkdtemp())
    data, out = tmp / "data", tmp / "out"
    data.mkdir()
    n3 = smart_planner(data / "sp3.xlsx", 3)
    n6 = smart_planner(data / "sp6.xlsx", 6, tie=True)
    smart_planner(data / "descuadre.xlsx", 3, skew_tp="TP1")
    smart_planner(data / "ancho.xlsx", 4, extra_col=True)
    spectrum(data / "alloc.xlsx", "Allocation Plan 1 Net Reach", [("TV", 100.0), ("Radio", 50.0)])
    spectrum(data / "reach.xlsx", "Plan reach (Department Stores, Plan 1, Personas 25-45 ABC+C, 4 weeks)",
             [("Net campaign reach", 0.8), ("TV", 0.6), ("Radio", 0.4)])
    wb = openpyxl.Workbook()
    wb.active["A1"] = "cualquier cosa"
    wb.save(data / "otro.xlsx")
    assert p.main([str(data), "--out", str(out)]) == 0

    for key, (names, inv, grps), n in [("sp3.xlsx", n3, 3), ("sp6.xlsx", n6, 6)]:
        m = load(out, key)
        lv = levels(m)
        assert m["source"]["herramienta"] == "smart_planner", key
        assert "error" not in lv.values() and "warning" not in lv.values(), (key, [v for v in m["validaciones"] if v["nivel"] != "ok"])
        ex = m["exercise"]
        assert (ex["periodo_inicio"], ex["periodo_fin"], ex["edad_min"], ex["edad_max"]) == ("2026-10-01", "2026-10-31", 18, 35)
        assert ex["target_descripcion"] is None and ex["universo"] == U and ex["industria"] is None
        assert [t["nombre"] for t in m["touchpoints"]] == names and m["touchpoints"][0]["curva"]["param_a"] == 10
        esc = m["escenarios"]
        assert [e["factor_presupuesto"] for e in esc] == [x / 10 for x in range(5, 16)]
        base = esc[5]["por_touchpoint"]
        assert len(base) == n and {t["touchpoint"]: (t["inversion"], round(t["grps"], 6)) for t in base} == \
            {nm: (a, round(g, 6)) for nm, a, g in zip(names, inv, grps)}, key
        assert len(m["alcance_por_frecuencia"]) == n + 1

    lv = levels(load(out, "descuadre.xlsx"))
    assert lv["inversion_vs_total"] == "error" and lv["inversion_vs_presupuesto"] == "error"

    m = load(out, "ancho.xlsx")
    assert levels(m)["escenarios_ancho"] == "error" and m["escenarios"] is None

    m = load(out, "alloc.xlsx")
    assert (m["source"]["herramienta"], m["source"]["subtipo"]) == ("spectrum", "allocation")
    assert m["exercise"]["campana"] == "Plan 1 Net Reach" and m["exercise"]["semanas"] is None
    assert [(t["nombre"], t["inversion"]) for t in m["touchpoints"]] == [("TV", 100.0), ("Radio", 50.0)]
    assert levels(m)["titulo_spectrum"] == "warning" and m["escenarios"] is None and m["totales"]["inversion"] is None

    m = load(out, "reach.xlsx")
    ex = m["exercise"]
    assert m["source"]["subtipo"] == "plan_reach" and m["source"]["fecha_export"] == "2026-09-25"
    assert (ex["anunciante"], ex["campana"], ex["target"], ex["semanas"], ex["edad_min"]) == \
        ("Department Stores", "Plan 1", "Personas 25-45 ABC+C", 4, 25)
    assert m["totales"]["alcance_neto"] == 0.8 and [t["alcance"] for t in m["touchpoints"]] == [0.6, 0.4]
    assert levels(m) == {"titulo_spectrum": "ok", "alcance_neto_rango": "ok"}

    m = load(out, "otro.xlsx")
    assert m["source"]["herramienta"] == "unknown" and levels(m) == {"formato": "error"}

    index = out / "index.md"
    index.write_text(index.read_text() + "\nNota manual.\n")
    assert p.main([str(data), "--out", str(out)]) == 0
    text = index.read_text()
    assert text.count("<!-- canonical:sp3.xlsx -->") == 1 and "Nota manual." in text
    assert CLIENT not in text and "MEXICO" not in text and "TP0" not in text  # sin valores de cliente
    print("ok")


if __name__ == "__main__":
    test()
