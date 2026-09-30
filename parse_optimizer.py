#!/usr/bin/env python3
"""Convierte outputs de optimizador (Havas Smart Planner / Havas Spectrum) a un modelo canónico JSON.

Uso:
    python parse_optimizer.py [ruta_archivo_o_carpeta] [--out ./canonical]

Sin ruta procesa DATA_DIR (o $HAVAS_DATA_DIR). Dependencias: openpyxl + stdlib.
"""
import argparse
import datetime as dt
import json
import logging
import os
import re
import sys
from pathlib import Path
from urllib.parse import quote

import openpyxl

# Único lugar a cambiar al mover el proyecto a la nube privada (o exportar HAVAS_DATA_DIR).
DATA_DIR = os.environ.get("HAVAS_DATA_DIR", "/Users/martinfranciscocortesbuendia/Desktop/EdgeNET/havasData/dataEx")
OUT_DIR = Path(__file__).resolve().parent / "canonical"

EXTS = {".xlsx", ".xlsm"}
TOL_SUM = 0.005    # sumas de inversión / GRPs contra total y presupuesto
TOL_SHARE = 0.01   # shares suman 1
TOL_RATIO = 0.02   # relaciones derivadas (impactos, alcance*frecuencia, CPGRP, CPM)
TOL_REACH = 0.005  # Media Mix redondea el alcance a 2 decimales

EXERCISE_FIELDS = ["pais", "industria", "anunciante", "marca", "campana", "periodo_inicio", "periodo_fin",
                   "presupuesto", "moneda", "target", "target_descripcion", "genero", "edad_min", "edad_max",
                   "universo", "semanas"]
TP_FIELDS = ["nombre", "color", "inversion", "share_inversion", "grps", "impactos_miles", "share_impactos",
             "alcance", "frecuencia", "costo_por_punto_alcance", "cpgrp", "cpm", "cpgrp_input", "cpm_input",
             "restriccion_min", "restriccion_max", "curva"]
TOTAL_FIELDS = ["inversion", "grps", "impactos_miles", "alcance_neto", "frecuencia", "costo_por_punto_alcance",
                "cpgrp", "cpm"]
MMT_COLS = ["inversion", "share_inversion", "grps", "impactos_miles", "share_impactos", "alcance", "frecuencia",
            "costo_por_punto_alcance", "cpgrp", "cpm"]  # columnas B..K de Media Mix table
EX_KEYS = {"country": "pais", "industry": "industria", "advertiser": "anunciante", "brand": "marca",
           "campaign": "campana", "budget": "presupuesto", "currency": "moneda", "target": "target",
           "target description": "target_descripcion", "gender": "genero", "target size": "universo"}
BLOCKS = ["inversion", "grps", "impactos_miles", "share_impactos", "alcance"]  # orden en Budget Scenarios

log = logging.getLogger("parse_optimizer")


# ---------------------------------------------------------------- utilidades

def num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip().replace(",", "")
    if re.fullmatch(r"-?\d+", s):
        return int(s)
    try:
        return float(s)
    except ValueError:
        return None


def text(v):
    s = str(v).strip() if v is not None else ""
    return s or None


def cell(row, i):
    return row[i] if i < len(row) else None


def close(a, b, tol):
    return abs(a - b) <= tol * max(abs(a), abs(b)) + 1e-9


def rows_of(wb, name):
    """Filas no vacías de la hoja, o None si la hoja no existe."""
    if name not in wb.sheetnames:
        return None
    return [r for r in wb[name].iter_rows(values_only=True) if any(v not in (None, "") for v in r)]


def width(row):
    return max((i + 1 for i, v in enumerate(row) if v not in (None, "")), default=0)


class Validations(list):
    def add(self, regla, nivel, mensaje, detalle=None):
        self.append({"regla": regla, "nivel": nivel, "mensaje": mensaje, "detalle": detalle or []})

    def rule(self, regla, fails, ok_msg, nivel="error", fail_msg="caso(s) fuera de tolerancia"):
        """fails=None: no hay datos para validar, no se registra."""
        if fails is None:
            return
        if fails:
            self.add(regla, nivel, f"{len(fails)} {fail_msg}", fails)
        else:
            self.add(regla, "ok", ok_msg)


def empty_model(path, key):
    return {
        "source": {"origen": None, "archivo": key, "herramienta": None, "subtipo": None, "version_detectada": None,
                   "fecha_export": None, "parseado_en": dt.datetime.now().isoformat(timespec="seconds")},
        "exercise": dict.fromkeys(EXERCISE_FIELDS),
        "touchpoints": None,
        "totales": dict.fromkeys(TOTAL_FIELDS),
        "alcance_por_frecuencia": None,
        "duplicaciones": None,
        "curvas_estimadas": None,
        "datapoints": None,
        "escenarios": None,
        "validaciones": Validations(),
    }


class Touchpoints(dict):
    """nombre normalizado -> dict canónico; conserva el orden de aparición."""
    def get_tp(self, name):
        k = name.strip().lower()
        if k not in self:
            self[k] = dict.fromkeys(TP_FIELDS)
            self[k]["nombre"] = name.strip()
        return self[k]


# ---------------------------------------------------------------- detección

def detect(wb):
    names = wb.sheetnames
    if "Exercise" in names and "Media Mix table" in names:
        return "smart_planner", None
    if len(names) == 1:
        top = [r for r in wb.worksheets[0].iter_rows(max_row=3, values_only=True)]
        a1 = " ".join(str(v) for v in (top[0] if top else ()) if v is not None)
        if re.search(r"Application:\s*Havas Spectrum", a1, re.I) and len(top) == 3:
            title = " ".join(str(v) for v in top[2] if v is not None)
            title = re.sub(r"^\s*Title:\s*", "", title, flags=re.I)
            if re.match(r"allocation\b", title, re.I):
                return "spectrum", "allocation"
            if re.match(r"plan reach\b", title, re.I):
                return "spectrum", "plan_reach"
    return "unknown", None


# ---------------------------------------------------------------- Smart Planner

def parse_exercise(rows, m):
    ex, src = m["exercise"], m["source"]
    for r in rows:
        i = next((i for i, v in enumerate(r) if v not in (None, "")), None)
        if i is None or not isinstance(r[i], str) or not r[i].strip().endswith(":"):
            continue  # títulos como "SCOPE" (celda combinada)
        key, val = r[i].strip().rstrip(":").strip().lower(), cell(r, i + 1)
        if key == "campaign period":
            parts = [p.strip() for p in str(val or "").split(",")]
            if len(parts) == 2:
                ex["periodo_inicio"], ex["periodo_fin"] = parts
        elif key == "age":
            parts = [num(p) for p in str(val or "").split(",")]
            if len(parts) == 2:
                ex["edad_min"], ex["edad_max"] = parts
        elif key == "date":
            src["fecha_export"] = val.isoformat() if isinstance(val, dt.datetime) else text(val)
        elif key in EX_KEYS:
            f = EX_KEYS[key]
            ex[f] = num(val) if f in ("presupuesto", "universo") else text(val)


def parse_smart_planner(wb, m, V):
    src = m["source"]
    src.update(origen="Havas Smart Planner", herramienta="smart_planner")
    tps, sheets, extra = Touchpoints(), {}, {}

    def body(name):
        rows = rows_of(wb, name)
        if rows is None:
            V.add("hoja_presente", "error", f"falta la hoja '{name}'")
            return None
        return rows[1:]  # sin encabezado

    def mmt():
        rows = body("Media Mix table")
        if rows is None:
            return
        for r in rows:
            name = text(r[0])
            vals = {f: num(cell(r, i + 1)) for i, f in enumerate(MMT_COLS)}
            if name and name.lower() == "total":
                t = m["totales"]
                for f in ("inversion", "grps", "impactos_miles", "frecuencia", "costo_por_punto_alcance", "cpgrp", "cpm"):
                    t[f] = vals[f]
                t["alcance_neto"] = vals["alcance"]
            elif name:
                tps.get_tp(name).update(vals)
        sheets["Media Mix table"] = [t["nombre"] for t in tps.values()]

    def simple(name, fields, target=None):
        rows = body(name)
        if rows is None:
            return
        out = []
        for r in rows:
            if not text(r[0]):
                continue
            vals = {f: (text(cell(r, i + 1)) if f == "color" else num(cell(r, i + 1))) for i, f in enumerate(fields)}
            if target is None:
                tps.get_tp(text(r[0])).update(vals)
            out.append({"touchpoint": text(r[0]), **vals})
        sheets[name] = sorted({o["touchpoint"] for o in out}, key=str.lower)
        if target:
            m[target] = out

    def curve_params():
        rows = body("Curve Parameters")
        if rows is None:
            return
        for r in rows:
            if text(r[0]):
                tps.get_tp(text(r[0]))["curva"] = {f"param_{c}": num(cell(r, i + 1)) for i, c in enumerate("abc")}
        sheets["Curve Parameters"] = [text(r[0]) for r in rows if text(r[0])]

    def rf():
        rows = body("Reach & Frequency")
        if rows is None:
            return
        m["alcance_por_frecuencia"] = [{"touchpoint": text(r[0]), **{f"f{k}": num(cell(r, k)) for k in range(1, 11)}}
                                       for r in rows if text(r[0])]
        sheets["Reach & Frequency"] = [x["touchpoint"] for x in m["alcance_por_frecuencia"]
                                       if x["touchpoint"].lower() != "total"]

    def cross():
        rows = body("Cross Indices")
        if rows is None:
            return
        m["duplicaciones"] = [{"touchpoint_a": text(r[0]), "touchpoint_b": text(cell(r, 1)), "indice": num(cell(r, 2))}
                              for r in rows if text(r[0])]
        sheets["Cross Indices"] = sorted({x for d in m["duplicaciones"] for x in (d["touchpoint_a"], d["touchpoint_b"]) if x})

    def plan_cost():
        rows = rows_of(wb, "Plan Cost table")
        if rows is None:
            V.add("hoja_presente", "error", "falta la hoja 'Plan Cost table'")
            return
        extra["plan_cost"] = {str(r[0]).strip().lower(): num(cell(r, 1)) for r in rows if r[0] is not None}

    steps = [
        ("Exercise", lambda: parse_exercise(rows_of(wb, "Exercise"), m)),
        ("Media Mix table", mmt),
        ("TP Colors", lambda: simple("TP Colors", ["color"])),
        ("Costs", lambda: simple("Costs", ["cpgrp_input", "cpm_input"])),
        ("Restrictions", lambda: simple("Restrictions", ["restriccion_min", "restriccion_max"])),
        ("Curve Parameters", curve_params),
        ("Reach & Frequency", rf),
        ("Cross Indices", cross),
        ("Curves Estimated", lambda: simple("Curves Estimated", ["grps", "alcance"], "curvas_estimadas")),
        ("Datapoints", lambda: simple("Datapoints", ["grps", "alcance"], "datapoints")),
        ("Plan Cost table", plan_cost),
        ("Budget Scenarios", lambda: parse_scenarios(wb, m, tps, sheets, V)),
    ]
    for name, fn in steps:
        try:
            fn()
        except Exception as e:  # una hoja rota no tumba el archivo
            log.exception("  falló la hoja '%s'", name)
            V.add("hoja_parseable", "error", f"no se pudo parsear '{name}': {type(e).__name__}: {e}")
    m["touchpoints"] = list(tps.values()) or None
    validate_smart_planner(m, tps, sheets, extra, V)


def parse_scenarios(wb, m, tps, sheets, V):
    rows = rows_of(wb, "Budget Scenarios")
    if rows is None:
        V.add("hoja_presente", "error", "falta la hoja 'Budget Scenarios'")
        return
    mix = [t for t in tps.values() if t["inversion"] is not None]
    n = len(mix)
    w = max(width(r) for r in rows)
    if not n or w != 5 + 5 * n:
        V.add("escenarios_ancho", "error", f"ancho real {w} no cuadra con 5 + 5*N (N={n} touchpoints en Media Mix table "
                                           f"=> {5 + 5 * n}); no se parsean los escenarios")
        return
    V.add("escenarios_ancho", "ok", f"ancho {w} = 5 + 5*{n}")
    header, data = rows[0], [r for r in rows[1:] if num(r[0]) is not None]
    col = lambda b, j: b * (n + 1) + 1 + j  # columna del touchpoint j en el bloque b

    # Orden: la fila cuyo presupuesto total = presupuesto del Exercise se cruza contra Media Mix (inversión, GRPs, impactos).
    budget, order = m["exercise"]["presupuesto"], None
    ref = min(data, key=lambda r: abs(num(r[0]) - budget), default=None) if budget else None
    if ref is not None and close(num(ref[0]), budget, TOL_SUM):
        match = []
        for j in range(n):
            vec = [num(ref[col(b, j)]) for b in range(3)]
            cands = [t for t in mix if all(v is not None and t[f] is not None and close(v, t[f], 0.01)
                                           for v, f in zip(vec, BLOCKS[:3]))]
            match.append(cands[0]["nombre"] if len(cands) == 1 else None)
        if None not in match and len(set(match)) == n:
            order = match
    by_header = [re.sub(r"^[^_]*_", "", str(header[col(0, j)] or "")) for j in range(n)]
    header_ok = ({h.lower() for h in by_header} == {t["nombre"].lower() for t in mix}
                 and all(re.sub(r"^[^_]*_", "", str(header[col(b, j)] or "")) == by_header[j]
                         for b in range(5) for j in range(n)))
    if order:
        V.add("escenarios_orden", "ok", "orden de touchpoints resuelto cruzando valores con Media Mix table")
        if header_ok and [h.lower() for h in by_header] != [o.lower() for o in order]:
            V.add("escenarios_encabezados", "warning", "los encabezados no coinciden con el orden deducido por valores",
                  [f"{h} -> {o}" for h, o in zip(by_header, order) if h.lower() != o.lower()])
    else:
        order = by_header if header_ok else [t["nombre"] for t in mix]
        V.add("escenarios_orden", "warning", "no se pudo deducir el orden por valores; se usa el orden en que aparecen "
              + ("en los encabezados" if header_ok else "en Media Mix table (posicional, no verificado)"))
    sheets["Budget Scenarios"] = sorted(order, key=str.lower)

    m["escenarios"] = []
    for i, r in enumerate(data, 1):
        total = num(r[0])
        m["escenarios"].append({
            "indice": i,
            "factor_presupuesto": round(total / budget, 2) if budget else None,
            "presupuesto_total": total,
            "por_touchpoint": [{"touchpoint": name, **{f: num(r[col(b, j)]) for b, f in enumerate(BLOCKS)}}
                               for j, name in enumerate(order)],
            "totales": {"grps": num(r[n + 1]), "impactos_miles": num(r[2 * (n + 1)]), "alcance": num(r[4 * (n + 1)]),
                        "_share_total": num(r[3 * (n + 1)])},
        })


def validate_smart_planner(m, tps, sheets, extra, V):
    ex, tot, tl = m["exercise"], m["totales"], list(tps.values())
    have = lambda *fs: [t for t in tl if all(t[f] is not None for f in fs)]

    inv = have("inversion")
    s = sum(t["inversion"] for t in inv)
    V.rule("inversion_vs_total", None if not inv or tot["inversion"] is None else
           ([] if close(s, tot["inversion"], TOL_SUM) else [f"suma {s:,.2f} vs total {tot['inversion']:,.2f}"]),
           "la suma por touchpoint cuadra con el total")
    V.rule("inversion_vs_presupuesto", None if not inv or ex["presupuesto"] is None else
           ([] if close(s, ex["presupuesto"], TOL_SUM) else [f"suma {s:,.2f} vs presupuesto {ex['presupuesto']:,.2f}"]),
           "la suma por touchpoint cuadra con el presupuesto del Exercise")
    for f in ("share_inversion", "share_impactos"):
        sh = have(f)
        s2 = round(sum(t[f] for t in sh), 6)
        V.rule(f"{f}_suma_1", None if not sh else ([] if abs(s2 - 1) <= TOL_SHARE + 1e-9 else [f"suma {s2}"]),
               f"{f} suma 1")
    g = have("grps")
    sg = sum(t["grps"] for t in g)
    V.rule("grps_vs_total", None if not g or tot["grps"] is None else
           ([] if close(sg, tot["grps"], TOL_SUM) else [f"suma {sg:,.2f} vs total {tot['grps']:,.2f}"]),
           "la suma de GRPs cuadra con el total")

    u = ex["universo"]
    if u:
        rows = [(t["nombre"], t["grps"], t["impactos_miles"]) for t in have("grps", "impactos_miles")]
        if tot["grps"] is not None and tot["impactos_miles"] is not None:
            rows.append(("Total", tot["grps"], tot["impactos_miles"]))
        V.rule("impactos_vs_grps_universo", [f"{n}: {i:,.2f} vs esperado {g_ * u / 1e5:,.2f}" for n, g_, i in rows
                                             if not close(i, g_ * u / 1e5, TOL_RATIO)] if rows else None,
               "impactos_miles ≈ grps * universo / 100 / 1000", "warning")

    f1 = {x["touchpoint"].lower(): x["f1"] for x in m["alcance_por_frecuencia"] or [] if x["f1"] is not None}
    rf_rows = have("alcance", "frecuencia", "grps")
    # ponytail: se usa el 1+ de Reach & Frequency (4 decimales) cuando existe; Media Mix redondea a 2
    V.rule("alcance_x_frecuencia_vs_grps", [f"{t['nombre']}: {f1.get(t['nombre'].lower(), t['alcance']) * t['frecuencia']:.4f}"
                                            f" vs {t['grps'] / 100:.4f}" for t in rf_rows
                                            if not close(f1.get(t["nombre"].lower(), t["alcance"]) * t["frecuencia"],
                                                         t["grps"] / 100, TOL_RATIO)] if rf_rows else None,
           "alcance * frecuencia ≈ grps / 100", "warning")

    al = [t["alcance"] for t in have("alcance")]
    net = tot["alcance_neto"]
    V.rule("alcance_neto_rango", None if net is None or not al else
           ([] if max(al) - TOL_REACH <= net <= sum(al) + TOL_REACH else [f"neto {net} fuera de [{max(al)}, {sum(al):.4f}]"]),
           "el alcance neto está entre el mayor individual y la suma")

    if m["alcance_por_frecuencia"]:
        V.rule("rf_decreciente", [x["touchpoint"] for x in m["alcance_por_frecuencia"]
                                  if any(a is not None and b is not None and b > a + 1e-9
                                         for a, b in zip(*[[x[f"f{k}"] for k in r] for r in (range(1, 10), range(2, 11))]))],
               "todas las series 1+ a 10+ son decrecientes")
        pairs = [(t["nombre"], t["alcance"], f1[t["nombre"].lower()]) for t in have("alcance") if t["nombre"].lower() in f1]
        if net is not None and "total" in f1:
            pairs.append(("Total", net, f1["total"]))
        V.rule("alcance_mmt_vs_rf_1mas", [f"{n}: {a} vs {b}" for n, a, b in pairs if abs(a - b) > TOL_REACH + 1e-9]
               if pairs else None, "el alcance de Media Mix coincide con 1+ de Reach & Frequency")

    rows = [(t["nombre"], t["inversion"], t["grps"], t["impactos_miles"], t["cpgrp"], t["cpm"]) for t in tl]
    rows.append(("Total", tot["inversion"], tot["grps"], tot["impactos_miles"], tot["cpgrp"], tot["cpm"]))
    cp = [f"{n}: CPGRP {c:,.2f} vs {i / g_:,.2f}" for n, i, g_, _, c, _ in rows
          if None not in (i, g_, c) and g_ and not close(c, i / g_, TOL_RATIO)]
    cm = [f"{n}: CPM {c:,.2f} vs {i / im:,.2f}" for n, i, _, im, _, c in rows
          if None not in (i, im, c) and im and not close(c, i / im, TOL_RATIO)]
    V.rule("cpgrp_vs_inversion_grps", cp if any(None not in (r[1], r[2], r[4]) for r in rows) else None,
           "CPGRP ≈ inversión / GRPs", "warning")
    V.rule("cpm_vs_inversion_impactos", cm if any(None not in (r[1], r[3], r[5]) for r in rows) else None,
           "CPM ≈ inversión / (impactos_miles * 1000) * 1000", "warning")

    pc = extra.get("plan_cost")
    if pc:
        checks = [("costo_por_punto_alcance", "reach point"), ("cpgrp", "total cpgrp"), ("cpm", "total cpm")]
        V.rule("totales_vs_plan_cost", [f"{f}: {tot[f]} vs {pc.get(k)}" for f, k in checks
                                        if tot[f] is not None and pc.get(k) is not None and not close(tot[f], pc[k], 0.001)],
               "los totales coinciden con Plan Cost table")

    base = {t["nombre"].lower() for t in tl if t["inversion"] is not None}
    if base:
        missing, extra_tp = [], []
        for sheet, names in sheets.items():
            got = {x.lower() for x in names}
            missing += [f"{sheet}: falta {t}" for t in sorted(base - got)]
            extra_tp += [f"{sheet}: sobra {t}" for t in sorted(got - base)]
        V.rule("touchpoints_en_todas_las_hojas", missing, "cada touchpoint del mix aparece en todas las hojas")
        if extra_tp:
            V.add("touchpoints_fuera_del_mix", "warning", "hay touchpoints que no están en Media Mix table", extra_tp)

    forced = [f"{t['nombre']}: {t['restriccion_min']}" for t in have("restriccion_min", "restriccion_max")
              if t["restriccion_min"] == t["restriccion_max"]]
    V.rule("restricciones_forzadas", forced if have("restriccion_min", "restriccion_max") else None,
           "ninguna restricción tiene mínimo = máximo", "warning")
    if forced:
        V[-1]["mensaje"] = f"{len(forced)} touchpoint(s) con mínimo = máximo: el mix venía forzado, el optimizador no optimizó"

    dup = [d["indice"] for d in m["duplicaciones"] or [] if d["indice"] is not None]
    if dup:
        if all(d == 100 for d in dup):
            V.add("duplicaciones_default", "warning", f"los {len(dup)} índices de duplicación valen 100 (valor por defecto)")
        else:
            V.add("duplicaciones_default", "ok", "hay índices de duplicación distintos de 100")

    esc = m["escenarios"]
    if esc:
        bad = [f"escenario {e['indice']}: {e['totales']['_share_total']}" for e in esc
               if e["totales"]["_share_total"] is not None and abs(e["totales"]["_share_total"] - 1) > 1e-6]
        V.rule("escenarios_share_total_1", bad, "share total = 1 en todos los escenarios", "warning")
        facs = [e["factor_presupuesto"] for e in esc if e["factor_presupuesto"] is not None]
        V.rule("escenarios_factor_rango", [str(f) for f in facs if not 0.5 <= f <= 1.5] if facs else None,
               f"factores de presupuesto entre 0.5 y 1.5 ({min(facs, default=0)} a {max(facs, default=0)})", "warning")
        for e in esc:
            e["totales"].pop("_share_total")


# ---------------------------------------------------------------- Spectrum

def parse_spectrum(wb, subtipo, m, V):
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    src, ex = m["source"], m["exercise"]
    src.update(origen=text(cell(rows[0], 1)) or "Havas Spectrum", herramienta="spectrum", subtipo=subtipo)
    raw_date = text(cell(rows[1], 1)) if len(rows) > 1 else None
    try:
        src["fecha_export"] = dt.datetime.strptime(raw_date, "%d/%m/%Y").date().isoformat()
    except (TypeError, ValueError):
        src["fecha_export"] = raw_date

    title = re.sub(r"^\s*Title:\s*", "", " ".join(str(v) for v in rows[2] if v is not None), flags=re.I)
    rest = re.sub(r"^(allocation|plan reach)\s*", "", title, flags=re.I).strip()
    inner = re.fullmatch(r"\((.*)\)", rest)
    if inner:
        parts = [p.strip() for p in inner.group(1).split(",")]
        weeks = next((p for p in parts if re.fullmatch(r"\d+\s*(weeks?|semanas?)", p, re.I)), None)
        if weeks:
            ex["semanas"] = int(re.match(r"\d+", weeks).group())
            parts.remove(weeks)
        if len(parts) == 3:  # categoría/anunciante, plan, target
            ex["anunciante"], ex["campana"], ex["target"] = parts
    elif rest:
        ex["campana"] = rest  # "Allocation Plan 1 Net Reach": solo trae el nombre del plan
    if ex["target"]:
        ages = re.search(r"\b(\d{1,2})\s*-\s*(\d{1,2})\b", ex["target"])
        if ages:
            ex["edad_min"], ex["edad_max"] = int(ages.group(1)), int(ages.group(2))
    nulls = [f for f in ("anunciante", "campana", "target", "semanas") if ex[f] is None]
    V.rule("titulo_spectrum", [f"sin {f}" for f in nulls], "título parseado completo", "warning",
           "campo(s) del título no identificados, quedan en null")

    tps = Touchpoints()
    for r in rows[5:]:  # datos desde la fila 6
        name, val = text(cell(r, 0)), num(cell(r, 1))
        if not name:
            continue
        if subtipo == "plan_reach" and name.lower() == "net campaign reach":
            m["totales"]["alcance_neto"] = val
        else:
            tps.get_tp(name)["alcance" if subtipo == "plan_reach" else "inversion"] = val
    m["touchpoints"] = list(tps.values()) or None

    al = [t["alcance"] for t in tps.values() if t["alcance"] is not None]
    net = m["totales"]["alcance_neto"]
    V.rule("alcance_neto_rango", None if net is None or not al else
           ([] if max(al) <= net <= sum(al) else [f"neto {net} fuera de [{max(al)}, {sum(al):.4f}]"]),
           "el alcance neto está entre el mayor individual y la suma")
    if subtipo == "plan_reach" and net is None:
        V.add("alcance_neto_presente", "error", "no se encontró 'Net campaign reach'")


# ---------------------------------------------------------------- reportes

def fmt(v):
    if v is None:
        return "—"
    if isinstance(v, float):
        return f"{v:,.4f}" if abs(v) < 10 else f"{v:,.2f}"
    if isinstance(v, int):
        return f"{v:,}"
    return str(v).replace("|", "\\|")


LIGHT = {"ok": "🟢", "warning": "🟡", "error": "🔴"}


def counts(m):
    c = {"ok": 0, "warning": 0, "error": 0}
    for v in m["validaciones"]:
        c[v["nivel"]] += 1
    return c


def render_md(m):
    s = m["source"]
    L = [f"# Modelo canónico: {s['archivo']}", "",
         f"- Herramienta: {s['herramienta']}" + (f" ({s['subtipo']})" if s["subtipo"] else "") + f" · Origen: {fmt(s['origen'])}",
         f"- Exportado: {fmt(s['fecha_export'])} · Parseado: {s['parseado_en']}", "",
         "## Ejercicio", "", "| Campo | Valor |", "|---|---|"]
    L += [f"| {k} | {fmt(v)} |" for k, v in m["exercise"].items()]
    L += ["", "## Mix de medios", ""]
    if m["touchpoints"]:
        cols = ["inversion", "share_inversion", "grps", "impactos_miles", "share_impactos", "alcance", "frecuencia",
                "costo_por_punto_alcance", "cpgrp", "cpm"]
        L += ["| Touchpoint | " + " | ".join(cols) + " |", "|---|" + "--:|" * len(cols)]
        L += [f"| {fmt(t['nombre'])} | " + " | ".join(fmt(t[c]) for c in cols) + " |" for t in m["touchpoints"]]
    else:
        L.append("Sin touchpoints.")
    L += ["", "## Totales", "", "| Campo | Valor |", "|---|---|"]
    L += [f"| {k} | {fmt(v)} |" for k, v in m["totales"].items()]
    c = counts(m)
    L += ["", "## Validaciones", "", f"🟢 {c['ok']} · 🟡 {c['warning']} · 🔴 {c['error']}", ""]
    for v in sorted(m["validaciones"], key=lambda v: ["error", "warning", "ok"].index(v["nivel"])):
        L.append(f"- {LIGHT[v['nivel']]} **{v['regla']}**: {v['mensaje']}")
        L += [f"  - {fmt(d)}" for d in v["detalle"][:10]]
        if len(v["detalle"]) > 10:
            L.append(f"  - … {len(v['detalle']) - 10} más (ver JSON)")
    return "\n".join(L) + "\n"


def index_section(m, key):
    """Sin valores de cliente: solo nombres de campo y conteos."""
    s, c = m["source"], counts(m)
    n = lambda x: len(x) if x else 0
    nulls = [f"exercise.{k}" for k, v in m["exercise"].items() if v is None] + \
            [f"totales.{k}" for k, v in m["totales"].items() if v is None]
    failing = sorted({v["regla"] for v in m["validaciones"] if v["nivel"] != "ok"})
    return "\n".join([
        f"## {key}", "",
        f"- Fecha: {s['parseado_en']}",
        f"- Formato: {s['herramienta']}" + (f" ({s['subtipo']})" if s["subtipo"] else ""),
        f"- Conteos: touchpoints {n(m['touchpoints'])} · escenarios {n(m['escenarios'])} · "
        f"filas R&F {n(m['alcance_por_frecuencia'])} · duplicaciones {n(m['duplicaciones'])} · "
        f"curvas {n(m['curvas_estimadas'])} · datapoints {n(m['datapoints'])}",
        f"- Campos nulos ({len(nulls)}): {', '.join(nulls) or '—'}",
        f"- Validaciones: 🟢 {c['ok']} · 🟡 {c['warning']} · 🔴 {c['error']}"
        + (f" — reglas con hallazgos: {', '.join(failing)}" if failing else ""),
        f"- Reportes: [JSON]({quote(key + '.canonical.json')}) · [Markdown]({quote(key + '.canonical.md')})",
    ])


def update_index(out, key, section):
    """Anexa o reemplaza solo la sección de este archivo; el resto del index queda intacto."""
    index = out / "index.md"
    txt = index.read_text(encoding="utf-8") if index.exists() else \
        "# Modelos canónicos\n\nUna sección por archivo procesado; se reemplaza al volver a procesarlo.\n"
    start, end = f"<!-- canonical:{key} -->", f"<!-- /canonical:{key} -->"
    block = f"{start}\n{section}\n{end}"
    pat = re.compile(re.escape(start) + r".*?" + re.escape(end), re.S)
    txt = pat.sub(lambda _: block, txt, count=1) if pat.search(txt) else txt.rstrip("\n") + "\n\n" + block + "\n"
    tmp = index.with_name("index.md.tmp")
    tmp.write_text(txt, encoding="utf-8")
    os.replace(tmp, index)


def process(path, key, out):
    m = empty_model(path, key)
    V = m["validaciones"]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        kind, subtipo = detect(wb)
        log.info("  formato: %s%s", kind, f" ({subtipo})" if subtipo else "")
        if kind == "smart_planner":
            parse_smart_planner(wb, m, V)
        elif kind == "spectrum":
            parse_spectrum(wb, subtipo, m, V)
        else:
            m["source"]["herramienta"] = "unknown"
            V.add("formato", "error", "formato no reconocido (ni Smart Planner ni Spectrum)")
    finally:
        wb.close()
    with open(out / f"{key}.canonical.json", "w", encoding="utf-8") as fh:
        json.dump(m, fh, ensure_ascii=False, indent=2, default=str)
    (out / f"{key}.canonical.md").write_text(render_md(m), encoding="utf-8")
    update_index(out, key, index_section(m, key))
    return m


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("path", nargs="?", default=DATA_DIR, help=f"archivo o carpeta (default: {DATA_DIR})")
    ap.add_argument("--out", default=OUT_DIR, help=f"carpeta de salida (default: {OUT_DIR})")
    args = ap.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", datefmt="%H:%M:%S")

    root, out = Path(args.path).expanduser().resolve(), Path(args.out).expanduser().resolve()
    if not root.exists():
        log.error("No existe: %s", root)
        return 2
    files = [root] if root.is_file() else sorted(
        p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in EXTS and not p.name.startswith("~$"))
    if not files:
        log.warning("No hay archivos .xlsx/.xlsm en %s", root)
        return 0
    out.mkdir(parents=True, exist_ok=True)

    failed = 0
    for i, f in enumerate(files, 1):
        key = f.name if root.is_file() else "__".join(f.relative_to(root).parts)
        log.info("[%d/%d] %s", i, len(files), f)
        try:
            c = counts(process(f, key, out))
            log.info("  validaciones: %d ok, %d warning, %d error", c["ok"], c["warning"], c["error"])
        except Exception:
            failed += 1
            log.exception("  no se pudo procesar %s", f)
    log.info("Listo: %d archivos, %d con error. Salida en %s", len(files), failed, out)
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
