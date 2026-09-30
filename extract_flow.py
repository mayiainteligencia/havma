#!/usr/bin/env python3
"""Extrae catálogo y reglas de negocio de flowcharts de planeación (flow anual y flow de campaña).

Uso:
    python extract_flow.py [ruta_archivo_o_carpeta] [--out ./rules]

Sin ruta procesa DATA_DIR (o $HAVAS_DATA_DIR). Dependencias: openpyxl + stdlib.
"""
import argparse
import datetime as dt
import json
import logging
import os
import re
import sys
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from urllib.parse import quote

import openpyxl
from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter, range_boundaries

# Único lugar a cambiar al mover el proyecto a la nube privada (o exportar HAVAS_DATA_DIR).
DATA_DIR = os.environ.get("HAVAS_DATA_DIR", "/Users/martinfranciscocortesbuendia/Desktop/EdgeNET/havasData/dataEx")
OUT_DIR = Path(__file__).resolve().parent / "rules"

EXTS = {".xlsx", ".xlsm"}
TOL = 0.01          # TRPs vs serie semanal y promedio vs TRPs/semanas
MAX_LIST = 500      # tope de filas detalladas en listas largas; los conteos son completos
ERRORS = {"#REF!", "#VALUE!", "#N/A", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!"}
MONTHS = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE",
          "NOVIEMBRE", "DICIEMBRE"]
DATE_RE = re.compile(r"(?<!\d)(\d{1,2})\s*[-/ ]?\s*(ENE|FEB|MAR|ABR|MAY|JUN|JUL|AGO|SEP|OCT|NOV|DIC)[A-Z]*\b")
MONTH_RE = re.compile(r"\b(" + "|".join(MONTHS) + r")\b")
PCT_RE = re.compile(r"(\d+(?:[.,]\d+)?)\s*%")
LEVEL_TOKENS = {"MMTO": "MANT", "MANT": "MANT", "MTTO": "MANT", "LANZ": "LANZ", "MIN": "MIN"}
MEDIA_WORDS = {"TV", "TVA", "TVP", "RADIO", "OOH", "DOOH", "DIGITAL", "CINE", "PRINT", "PRENSA", "REVISTAS", "CTV"}
NOT_MEDIA = re.compile(r"^(INVERSI|TOTAL|GRAN TOTAL|MEDIOS|PRESUPUESTO|DIFERENCIA)")
BUDGET_RE = re.compile(r"PPTO|RECORTE|BAJA|ADICIONAL|ELIMIN")
CUT_RE = re.compile(r"RECORTE|BAJA|ELIMIN")
NOT_BUDGET_LINE = re.compile(r"^(TOTAL|MAS |EN FLOW)")  # totales de la hoja, no líneas
BASE_RE = re.compile(r"^PPTO (DE )?MKT")
CAL_STEP = re.compile(r"^=\+?\$?[A-Z]{1,3}\$?(\d+)\+7$")  # =+G47+7: fechas del calendario

log = logging.getLogger("extract_flow")
_merge_warned = False


# ---------------------------------------------------------------- utilidades

def key(v):
    """Normaliza para comparar: mayúsculas, sin acentos, espacios colapsados."""
    s = unicodedata.normalize("NFKD", str(v)) if v is not None else ""
    return " ".join("".join(ch for ch in s if not unicodedata.combining(ch)).upper().split())


def txt(v):
    if not isinstance(v, str):
        return None
    s = v.strip()
    return s if s and s not in ERRORS and not s.startswith("=") else None


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def close(a, b, tol=TOL):
    return abs(a - b) <= tol * max(abs(a), abs(b)) + 1e-9


def cell(row, c):
    return row[c] if c is not None and c < len(row) else None


def col(c):
    return get_column_letter(c + 1)


def merged_owner(path, ws):
    """{(fila0, col0): (fila0, col0) de la celda superior izquierda} para rangos combinados de una sola columna."""
    global _merge_warned
    owner = {}
    try:  # read_only no expone merged_cells; se lee el XML (_worksheet_path es privado de openpyxl 3.1)
        with zipfile.ZipFile(path) as z, z.open(ws._worksheet_path) as fh:
            for _, el in ET.iterparse(fh):
                if el.tag.endswith("}mergeCell"):
                    c1, r1, c2, r2 = range_boundaries(el.get("ref"))
                    if c1 == c2:
                        for r in range(r1, r2 + 1):
                            owner[(r - 1, c1 - 1)] = (r1 - 1, c1 - 1)
                elif el.tag.endswith("}row"):
                    el.clear()
    except Exception as e:
        if not _merge_warned:
            log.warning("Celdas combinadas no disponibles, se usa relleno hacia abajo (%s: %s)", type(e).__name__, e)
            _merge_warned = True
    return owner


def find_sheet(names, pred):
    return next((n for n in names if pred(key(n))), None)


# ---------------------------------------------------------------- fórmulas

def constants_in(toks):
    """Constantes junto a + - * / ^ (o fórmula solo numérica). Devuelve [(valor, operación, es_base)].

    es_base: primer número de una fórmula sin referencias (=128*1.04): es un dato capturado, no un factor."""
    arith = lambda x: x is not None and x.type == "OPERATOR-INFIX" and x.value in ("+", "-", "*", "/", "^")
    only = all(t.subtype == "NUMBER" for t in toks if t.type == "OPERAND") and not any(t.type == "FUNC" for t in toks)
    out, seen_number = [], False
    for i, t in enumerate(toks):
        if t.type != "OPERAND" or t.subtype != "NUMBER":
            continue
        j = i - 1
        while j >= 0 and toks[j].type == "OPERATOR-PREFIX":
            j -= 1
        prev = toks[j] if j >= 0 else None
        nxt = toks[i + 1] if i + 1 < len(toks) else None
        pct = nxt is not None and nxt.type == "OPERATOR-POSTFIX" and nxt.value == "%"
        after = (toks[i + 2] if i + 2 < len(toks) else None) if pct else nxt
        if not (only or arith(prev) or arith(after)):
            continue  # argumentos de función: SUBTOTAL(9, ...), ROUND(x, 2)
        v = float(t.value) / (100 if pct else 1)
        if v in (0, 1):
            continue  # ponytail: 0 y 1 son contadores/ajustes triviales (=A1+1), ruido para el planner
        out.append((v, prev.value if arith(prev) else after.value if arith(after) else "=", only and not seen_number))
        seen_number = True
    return out


def scan_formulas(wb_f, dep_sheet=None):
    hard, n_hard, consts, deps, n_cal = [], 0, {}, defaultdict(Counter), 0
    for ws in wb_f.worksheets:
        for r, row in enumerate(ws.iter_rows(values_only=True), 1):
            for c, v in enumerate(row):
                f = getattr(v, "text", v)
                if not (isinstance(f, str) and f.startswith("=")):
                    continue
                m = CAL_STEP.match(f.replace(" ", ""))
                if m and int(m.group(1)) == r:
                    n_cal += 1  # incremento de semana del calendario, no es regla de negocio
                    continue
                try:
                    toks = [t for t in Tokenizer(f).items if t.type != "WHITE-SPACE"]
                except Exception:
                    continue
                where = f"{ws.title}!{col(c)}{r}"
                nums = constants_in(toks)
                if nums:
                    n_hard += 1
                    if len(hard) < MAX_LIST:
                        hard.append({"hoja": ws.title, "celda": f"{col(c)}{r}", "formula": f,
                                     "numeros": [f"{op}{fv(v)}" for v, op, _ in nums]})
                for v, op, base in nums:
                    if base:
                        continue
                    k = f"{v:g}"
                    d = consts.setdefault(k, {"valor": v, "frecuencia": 0, "operaciones": Counter(), "hojas": Counter(),
                                              "celdas": [], "ejemplo": f})
                    d["frecuencia"] += 1
                    d["operaciones"][op] += 1
                    d["hojas"][ws.title] += 1
                    if len(d["celdas"]) < 20:
                        d["celdas"].append(where)
                if ws.title == dep_sheet:
                    for t in toks:
                        if t.subtype == "RANGE" and "!" in t.value:
                            sheet = t.value.rsplit("!", 1)[0]
                            deps[r][sheet[1:-1].replace("''", "'") if sheet.startswith("'") else sheet] += 1
    # recurrentes, más multiplicadores tipo porcentaje (*66%, *1.04) aunque aparezcan una sola vez
    recurrent = sorted((d for d in consts.values() if d["frecuencia"] >= 2 or ("*" in d["operaciones"] and 0 < d["valor"] < 2)),
                       key=lambda d: -d["frecuencia"])
    for d in recurrent:
        d["interpretacion"] = interpret(d["valor"], d["operaciones"].most_common(1)[0][0])
        d["operaciones"], d["hojas"] = dict(d["operaciones"]), dict(d["hojas"])
    return ({"total": n_hard, "formulas": hard, "omitidas_calendario_+7": n_cal}, recurrent, deps)


def fv(v):
    """1000000.0 -> '1,000,000'; 1.04 -> '1.04'."""
    return f"{int(v):,}" if float(v).is_integer() else f"{v:g}"


def interpret(v, op):
    """Lectura tentativa de un factor, redactada para que el planner la confirme."""
    if v in (100, 1000, 1_000_000):
        return f"conversión de unidades ({op}{fv(v)})"
    if op == "*" and 0 < v < 0.2:
        return f"aplica {v * 100:g}% (¿comisión o fee?)"
    if op == "*" and 0 < v < 1:
        return f"toma el {v * 100:g}% (equivale a descontar {100 - v * 100:g}%)"
    if op == "*" and 1 < v < 2:
        return f"incremento de {(v - 1) * 100:g}%"
    if op == "/" and v == 7:
        return "reparto entre 7 (¿días de la semana?)"
    if op == "/" and v in (4, 12, 52):
        return f"reparto entre {fv(v)} (¿{ {4: 'semanas del mes', 12: 'meses del año', 52: 'semanas del año'}[v]}?)"
    if op == "/":
        return f"reparto entre {fv(v)}"
    if op in ("+", "-"):
        return f"ajuste fijo de {op}{fv(v)}"
    return f"factor {op}{fv(v)}"


def text_rules(grids):
    """Porcentajes declarados en textos (encabezados y notas): 'Inflación 5%', 'DESC 20%', 'se considera 4%'."""
    rules = {}
    for sheet, g in grids.items():
        for r, row in enumerate(g):
            for c, v in enumerate(row):
                s = txt(v)
                if s and PCT_RE.search(s):
                    d = rules.setdefault(key(s), {"texto": " ".join(s.split()), "porcentajes": sorted(set(PCT_RE.findall(s))),
                                                  "frecuencia": 0, "celdas": []})
                    d["frecuencia"] += 1
                    if len(d["celdas"]) < 10:
                        d["celdas"].append(f"{sheet}!{col(c)}{r + 1}")
    return sorted(rules.values(), key=lambda d: -d["frecuencia"])


# ---------------------------------------------------------------- flow anual

def calendar(g):
    """Primera fila con >= 4 fechas consecutivas a +7 días, más las filas de meses y trimestres de encima."""
    for r, row in enumerate(g):
        best, run = [], []
        for c, v in enumerate(row):
            d = v if isinstance(v, dt.datetime) else None
            if d and run and run[-1][0] == c - 1 and (d - run[-1][1]).days == 7:
                run.append((c, d))
            else:
                run = [(c, d)] if d else []
            best = run if len(run) > len(best) else best
        if len(best) >= 4:
            cols = [c for c, _ in best]
            labels = {}
            for name, pred in (("mes", lambda s: s in MONTHS),
                               ("trimestre", lambda s: re.fullmatch(r"[1-4]Q|Q[1-4]|[1-4]T|T[1-4]|TRIMESTRE \d", s))):
                for rr in range(r - 1, max(r - 6, -1), -1):
                    hits = {c: key(g[rr][c]) for c in cols if pred(key(cell(g[rr], c)))}
                    if len(hits) >= 2 or (hits and name == "trimestre"):
                        labels[name] = (rr, hits)
                        break
            weeks = []
            for i, (c, d) in enumerate(best, 1):
                w = {"semana": i, "fecha": d.date().isoformat(), "columna": col(c)}
                for name, (_, hits) in labels.items():
                    w[name] = next((hits[k] for k in sorted(hits, reverse=True) if k <= c), None)
                weeks.append(w)
            info = {"fila": r + 1, "fila_meses": labels["mes"][0] + 1 if "mes" in labels else None,
                    "fila_trimestres": labels["trimestre"][0] + 1 if "trimestre" in labels else None,
                    "semanas": weeks}
            return info, cols
    return None, []


HEADERS = [
    ("categoria", lambda h: h == "CATEGORIA"), ("campana", lambda h: h.startswith("CAMPANA")),
    ("target", lambda h: h == "TARGET"), ("medios", lambda h: h == "MEDIOS"),
    ("trps", lambda h: h.startswith("TRP") and "TOTAL" in h), ("promedio", lambda h: h in ("WKS PROM SEM", "WEEKLYS")),
    ("peso", lambda h: h == "PESO CAMPANA"), ("semanas", lambda h: h in ("SEMANAS AL AIRE", "WOA")),
    ("lanz", lambda h: h == "LANZAMIENTO"), ("mant", lambda h: h == "MANTENIMIENTO"), ("min", lambda h: h == "MINIMO"),
]


def norm_pattern(p):
    toks = [key(t) for t in re.split(r"[|/]", p) if t.strip()]
    return " | ".join(LEVEL_TOKENS.get(t, t) for t in toks)


def parse_annual(path, grids, wb_v):
    medios, hier, flights, levels, tbd = {}, {}, [], {}, 0
    cal_info, _ = calendar(grids["FLOW"])
    for name in [n for n in grids if n == "FLOW" or n.startswith("Flow ")]:
        g = grids[name]
        owner = merged_owner(path, wb_v[name])
        _, week_cols = calendar(g)
        heads = [r for r, row in enumerate(g) if {"CATEGORIA", "MEDIOS"} <= {key(v) for v in row if isinstance(v, str)}]
        for i, h in enumerate(heads):
            cm = {}
            for c, v in enumerate(g[h]):
                for f, pred in HEADERS:
                    if f not in cm and isinstance(v, str) and pred(key(v)):
                        cm[f] = c
            end = heads[i + 1] if i + 1 < len(heads) else len(g)
            cur = {"categoria": None, "campana": None, "target": None}
            for r in range(h + 1, end):
                row = g[r]
                first = [key(v) for v in row[:6] if isinstance(v, str)]
                if any(s.startswith("GRAN TOTAL") for s in first):
                    break
                if key(cell(row, cm.get("categoria"))).startswith("TOTAL"):
                    cur = dict.fromkeys(cur)
                    continue
                for f in cur:
                    c = cm.get(f)
                    if c is None:
                        continue
                    own = owner.get((r, c))
                    v = txt(cell(g[own[0]], own[1])) if own else txt(cell(row, c))
                    if v:
                        cur[f] = v
                    elif f == "campana" and cm.get("categoria") is not None and txt(cell(row, cm["categoria"])):
                        cur[f] = None  # nueva categoría sin campaña en la misma fila
                medio = txt(cell(row, cm.get("medios")))
                camp = cur["campana"]
                if medio and not NOT_MEDIA.match(key(medio)):
                    m = medios.setdefault(key(medio), {"medio": key(medio), "variantes": Counter(), "hojas": Counter()})
                    m["variantes"][" ".join(medio.split())] += 1
                    m["hojas"][name] += 1
                    if camp and key(camp).startswith("TBD"):
                        tbd += 1
                    elif camp:
                        k = (cur["categoria"], camp, cur["target"])
                        hh = hier.setdefault(k, {"categoria": k[0], "campana": k[1], "target": k[2], "hojas": set(), "filas": 0})
                        hh["hojas"].add(name)
                        hh["filas"] += 1
                if not camp:
                    continue
                lv = [num(cell(row, cm.get(f))) for f in ("lanz", "mant", "min")]
                if any(x is not None for x in lv) and key(camp) not in levels:
                    levels[key(camp)] = {"campana": camp, "hoja": name, "fila": r + 1,
                                         **dict(zip(("lanzamiento", "mantenimiento", "minimo"), lv))}
                peso, trps = txt(cell(row, cm.get("peso"))), num(cell(row, cm.get("trps")))
                if peso or trps is not None:
                    serie = [x for x in (num(cell(row, c)) for c in week_cols) if x is not None]
                    flights.append({
                        "hoja": name, "fila": r + 1, "categoria": cur["categoria"], "campana": camp, "medio": medio,
                        "patron": peso, "patron_normalizado": norm_pattern(peso) if peso else None,
                        "trps_totales": trps, "semanas_al_aire": num(cell(row, cm.get("semanas"))),
                        "promedio_semanal": num(cell(row, cm.get("promedio"))),
                        "suma_serie": sum(serie) if serie else None,
                    })

    bad_trps, bad_prom = [], []
    for f in flights:
        t, s, p, w = f["trps_totales"], f["suma_serie"], f["promedio_semanal"], f["semanas_al_aire"]
        where = f"{f['hoja']}!fila {f['fila']} ({f['campana']} / {f['medio']})"
        if t and s is not None and not close(t, s):
            bad_trps.append({"donde": where, "trps_totales": t, "suma_serie": s})
        if t and p is not None and w and not close(p, t / w):
            bad_prom.append({"donde": where, "promedio_semanal": p, "trps_entre_semanas": t / w})

    patterns = {}
    for f in flights:
        if not f["patron_normalizado"]:
            continue
        d = patterns.setdefault(f["patron_normalizado"], {"patron": f["patron_normalizado"], "variantes": Counter(),
                                                          "filas": 0, "campanas": set(), "medios": Counter()})
        d["variantes"][f["patron"].strip()] += 1
        d["filas"] += 1
        d["campanas"].add(key(f["campana"]))
        d["medios"][f["medio"]] += 1
    for d in patterns.values():
        d["niveles"] = [{k: levels[c][k] for k in ("lanzamiento", "mantenimiento", "minimo")}
                        for c in sorted(d["campanas"]) if c in levels]
        d["campanas"] = len(d["campanas"])

    return {
        "calendario": cal_info,
        "medios": sorted(({**m, "variantes": dict(m["variantes"]), "hojas": dict(m["hojas"])} for m in medios.values()),
                         key=lambda m: -sum(m["variantes"].values())),
        "jerarquia": {"combinaciones": [{**h, "hojas": sorted(h["hojas"])} for h in hier.values()],
                      "filas_placeholder_tbd": tbd},
        "flighting": {
            "registros": flights,
            "niveles_por_campana": list(levels.values()),
            "patrones": sorted(({**d, "variantes": dict(d["variantes"]), "medios": dict(d["medios"])}
                                for d in patterns.values()), key=lambda d: -d["campanas"]),
            "validaciones": {"trps_vs_serie_semanal": bad_trps, "promedio_vs_trps_entre_semanas": bad_prom},
        },
    }


def budget_adjustments(g, sheet):
    out = []
    for r, row in enumerate(g):
        for c, v in enumerate(row):
            s = txt(v)
            if not s or not BUDGET_RE.search(key(s)) or NOT_BUDGET_LINE.match(key(s)):
                continue
            amount = next((num(x) for x in row[c + 1:c + 4] if num(x) is not None), None)
            if amount is None:
                continue
            k = key(s)
            m = DATE_RE.search(k)
            fecha = f"{int(m.group(1))}-{m.group(2)}" if m else (MONTH_RE.search(k).group(1) if MONTH_RE.search(k) else None)
            out.append({"hoja": sheet, "celda": f"{col(c)}{r + 1}", "concepto": " ".join(s.split()),
                        "tipo": "presupuesto_base" if BASE_RE.match(k) else
                                "recorte" if CUT_RE.search(k) or amount < 0 else "adicional",
                        "monto": amount, "fecha_texto": fecha})
    return out


# ---------------------------------------------------------------- flow de campaña

def columns_catalog(g, fields, numbers=False):
    """fields: {nombre: {encabezados normalizados}}. Un renglón con >= 2 encabezados conocidos redefine columnas."""
    known = set().union(*fields.values())
    out, cm = {f: Counter() for f in fields}, None
    for row in g:
        keys = [key(v) if isinstance(v, str) else "" for v in row]
        if sum(k in known for k in keys) >= 2:
            cm = {f: [c for c, k in enumerate(keys) if k in hs] for f, hs in fields.items()}
            continue
        if cm:
            for f, cols in cm.items():
                for c in cols:
                    v = txt(cell(row, c)) or (f"{cell(row, c):g}" if numbers and num(cell(row, c)) is not None else None)
                    if v and v != "[REDACTED]":
                        out[f][" ".join(v.split())] += 1
    return {f: dict(c.most_common()) for f, c in out.items()}


def sumario(g):
    h = next((r for r, row in enumerate(g) if "PROVEEDOR" in {key(v) for v in row if isinstance(v, str)}), None)
    if h is None:
        return None
    keys = [key(v) if isinstance(v, str) else "" for v in g[h]]
    cm = {f: next((c for c, k in enumerate(keys) if pred(k)), None) for f, pred in
          (("medio", lambda k: k.startswith("MEDIO")), ("proveedor", lambda k: k == "PROVEEDOR"),
           ("razon", lambda k: k.startswith("RAZON SOCIAL")))}
    body = g[h + 1:]
    tipo_cols = []
    for c in range((cm["razon"] or cm["proveedor"] or 0) + 1, max(map(len, body), default=0)):
        vals = [key(txt(cell(r, c))) for r in body if txt(cell(r, c))]
        if len(vals) >= 2 and sum(bool(re.fullmatch(r"[A-Z]{2,10}(/[A-Z]{2,10})*", v)) for v in vals) >= 0.8 * len(vals):
            tipo_cols.append(c)
    provs, tipos, medio = {}, Counter(), None
    for row in body:
        medio = txt(cell(row, cm["medio"])) or medio
        p, rs = txt(cell(row, cm["proveedor"])), txt(cell(row, cm["razon"]))
        if p or rs:
            provs.setdefault((key(medio), p, rs), {"medio": " ".join((medio or "").split()) or None, "proveedor": p,
                                                   "razon_social": rs})
        for c in tipo_cols:
            if txt(cell(row, c)) and re.fullmatch(r"[A-Z]{2,10}(/[A-Z]{2,10})*", key(cell(row, c))):
                tipos[key(cell(row, c))] += 1
    return {"proveedores": list(provs.values()), "tipos_de_compra": dict(tipos.most_common()),
            "columnas_tipo_compra": [col(c) for c in tipo_cols]}


def bitacora(g, sheet):
    h = next((r for r, row in enumerate(g) if "FECHA" in {key(v) for v in row if isinstance(v, str)}), None)
    head = []
    for row in g[:h if h is not None else len(g)]:
        c = next((c for c, v in enumerate(row) if txt(v)), None)
        if c is not None:
            head.append({"campo": txt(row[c]).rstrip(":").strip(), "lleno": any(v not in (None, "") and str(v).strip()
                                                                                  for v in row[c + 1:])})
    if h is None:
        return {"hoja": sheet, "campos_encabezado": head, "fila_registro": None, "campos_registro": [], "entradas_llenas": 0}
    fields = [(c, txt(v)) for c, v in enumerate(g[h]) if txt(v)]
    entries = []
    for row in g[h + 1:]:
        e = {f: cell(row, c) for c, f in fields if cell(row, c) is not None and str(cell(row, c)).strip()}
        if e:
            entries.append(e)
    names = " ".join(key(f) for _, f in fields) + " " + " ".join(key(x["campo"]) for x in head)
    version = next((x for x in head if key(x["campo"]).startswith("VERSION")), None)
    return {
        "hoja": sheet, "campos_encabezado": head, "fila_registro": h + 1, "campos_registro": [f for _, f in fields],
        "campos_esperados": {k: bool(re.search(p, names)) for k, p in
                             (("version", r"VERSION"), ("fecha", r"FECHA"), ("quien_pidio", r"SOLICIT|PIDI"),
                              ("quien_envio", r"ENVI"), ("inversion", r"INVERSION"))},
        "version_llena": bool(version and version["lleno"]),
        "entradas_llenas": len(entries), "entradas": entries,
    }


def dependencies(g, deps):
    lc = next((c for row in g[:15] for c, v in enumerate(row) if key(v) == "MEDIA"), 1)
    starts = [r for r, row in enumerate(g) if (k := key(txt(cell(row, lc)))) and k.split()[0] in MEDIA_WORDS]
    out = []
    for i, s in enumerate(starts):
        e = starts[i + 1] if i + 1 < len(starts) else len(g)
        src = Counter()
        for r in range(s + 1, e + 1):  # deps usa filas 1-based
            src.update(deps.get(r, {}))
        out.append({"bloque": " ".join(txt(g[s][lc]).split()), "filas": f"{s + 1}-{e}",
                    "fuentes": dict(src.most_common()), "externas": [k for k in src if "[" in k]})
    return out


# ---------------------------------------------------------------- archivo

def detect(names):
    lower = {n.strip().lower() for n in names}
    if "FLOW" in names and any(n.startswith("Flow ") or n.startswith("Resumen Escenario") for n in names):
        return "flow_anual"
    if {"bitacora", "flow", "kpis"} <= lower:
        return "flow_campana"
    return "unknown"


def client_terms(rep, grids):
    """Nombres de cliente/marca/campaña/proveedor a ocultar en catalog.md."""
    terms = set()
    for g in grids.values():
        for row in g[:10]:
            for c, v in enumerate(row):
                if key(v).rstrip(":").strip() in ("CLIENTE", "CAMPANA", "MARCA"):
                    terms.add(txt(next((x for x in row[c + 1:] if txt(x)), None)))
    for h in (rep.get("jerarquia") or {}).get("combinaciones", []):
        terms.update((h["categoria"], h["campana"]))
    for n in grids:
        if n.startswith("Flow "):
            terms.add(n[5:].strip())
    for p in ((rep.get("catalogos") or {}).get("sumario") or {}).get("proveedores", []):
        terms.update((p["proveedor"], p["razon_social"]))
    return sorted(t for t in terms if t and len(t) >= 3 and key(t) not in ("TBD", "[REDACTED]"))


def process(path, key_, out):
    wb_f = openpyxl.load_workbook(path, read_only=True, data_only=False)
    wb_v = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        kind = detect(wb_v.sheetnames)
        log.info("  formato: %s", kind)
        rep = {"source": {"archivo": key_, "ruta": str(path), "formato": kind,
                          "procesado_en": dt.datetime.now().isoformat(timespec="seconds")}}
        if kind != "unknown":
            grids = {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb_v.worksheets}
            dep_sheet = find_sheet(grids, lambda k: k == "FLOW") if kind == "flow_campana" else None
            hard, consts, deps = scan_formulas(wb_f, dep_sheet)
            rep.update(formulas_hardcode=hard, constantes=consts, reglas_en_texto=text_rules(grids))
            if kind == "flow_anual":
                rep.update(parse_annual(path, grids, wb_v))
                extra = find_sheet(grids, lambda k: "EXTRA BUDGET" in k)
                rep["ajustes_presupuesto"] = budget_adjustments(grids[extra], extra) if extra else None
            else:
                cat = {}
                for label, pred, fields in [
                    ("radio", lambda k: k == "RADIO", {"plazas": {"PLAZA"}, "estaciones": {"SIGLAS"},
                                                       "proveedores": {"PROVEEDOR"}, "formatos": {"FORMATO_HM"},
                                                       "tipos": {"TIPO"}, "grupos": {"GRUPO"}}),
                    ("planta_ooh", lambda k: k.startswith("PLANTA"), {"clase": {"PLANTA"}, "tipos": {"TIPO"},
                                                                     "formatos": {"FORMATO"}, "tipos_de_compra": {"TIPO DE COMPRA"},
                                                                     "periodos": {"PERIODO"}}),
                    ("digital", lambda k: k == "DIGITAL", {"plataformas": {"PLATAFORMA"}, "tipos_de_pauta": {"TIPO DE PAUTA"},
                                                          "tipos_de_compra": {"TIPO DE COMPRA"}, "medios": {"MEDIO"},
                                                          "placements": {"PLACEMENT"}, "kpis": {"KPI"}}),
                ]:
                    s = find_sheet(grids, pred)
                    cat[label] = {"hoja": s, **columns_catalog(grids[s], fields)} if s else None
                s = find_sheet(grids, lambda k: k == "LISTAS")
                if s and grids[s]:
                    hdr = {key(v): {key(v)} for v in grids[s][0] if txt(v)}
                    cat["listas"] = {"hoja": s, **columns_catalog(grids[s], hdr, numbers=True)}
                else:
                    cat["listas"] = None
                s = find_sheet(grids, lambda k: k == "SUMARIO")
                cat["sumario"] = {"hoja": s, **sumario(grids[s])} if s and sumario(grids[s]) else None
                rep["catalogos"] = cat
                s = find_sheet(grids, lambda k: k == "BITACORA")
                rep["bitacora"] = bitacora(grids[s], s)
                rep["dependencias"] = dependencies(grids[dep_sheet], deps)
            rep["terminos_cliente"] = client_terms(rep, grids)
    finally:
        wb_f.close()
        wb_v.close()
    with open(out / f"{key_}.rules.json", "w", encoding="utf-8") as fh:
        json.dump(rep, fh, ensure_ascii=False, indent=2, default=str)
    (out / f"{key_}.rules.md").write_text(render_md(rep), encoding="utf-8")
    return rep


# ---------------------------------------------------------------- reportes

def md(v):
    if v is None:
        return "—"
    if isinstance(v, float):
        return f"{v:,.2f}"
    return str(v).replace("|", "\\|").replace("\n", " ")


def table(headers, rows, limit=None):
    rows = list(rows)
    L = ["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
    L += ["| " + " | ".join(md(x) for x in r) + " |" for r in rows[:limit]]
    if limit and len(rows) > limit:
        L.append(f"\n… {len(rows) - limit} más en el JSON.")
    return L if rows else ["Sin datos."]


def render_md(rep):
    s = rep["source"]
    L = [f"# Reglas: {s['archivo']}", "", f"- Formato: {s['formato']} · Procesado: {s['procesado_en']}", ""]
    if s["formato"] == "unknown":
        return "\n".join(L + ["No es un flow anual (FLOW + Flow*/Resumen Escenario) ni de campaña (Bitacora + Flow + Kpis)."]) + "\n"
    if s["formato"] == "flow_anual":
        cal = rep["calendario"]
        L += ["## 1. Calendario semanal", ""]
        if cal:
            w = cal["semanas"]
            L += [f"Fila de fechas {cal['fila']} (meses: {md(cal['fila_meses'])}, trimestres: {md(cal['fila_trimestres'])}); "
                  f"{len(w)} semanas, de {w[0]['fecha']} ({w[0]['columna']}) a {w[-1]['fecha']} ({w[-1]['columna']}).", ""]
            L += table(["Semana", "Fecha", "Columna", "Mes", "Trimestre"],
                       ((x["semana"], x["fecha"], x["columna"], x.get("mes"), x.get("trimestre")) for x in w))
        else:
            L.append("No se encontró una fila de fechas semanales.")
        L += ["", "## 2. Catálogo de medios", ""]
        L += table(["Medio normalizado", "Variantes (apariciones)", "Hojas"],
                   ((m["medio"], ", ".join(f"«{k}» ({v})" for k, v in m["variantes"].items()), len(m["hojas"]))
                    for m in rep["medios"]))
        j = rep["jerarquia"]
        L += ["", "## 3. Jerarquía categoría → campaña → target", "",
              f"{len(j['combinaciones'])} combinaciones; {j['filas_placeholder_tbd']} filas con campaña TBD omitidas.", ""]
        L += table(["Categoría", "Campaña", "Target", "Hojas", "Filas"],
                   ((h["categoria"], h["campana"], h["target"], ", ".join(h["hojas"]), h["filas"]) for h in j["combinaciones"]))
        fl = rep["flighting"]
        L += ["", "## 4. Reglas de flighting", "", "### Patrones de PESO CAMPAÑA", ""]
        L += table(["Patrón normalizado", "Campañas", "Filas", "Variantes escritas", "Niveles (lanz/mant/mín)"],
                   ((p["patron"], p["campanas"], p["filas"], ", ".join(f"«{k}»" for k in p["variantes"]),
                     "; ".join(f"{md(n['lanzamiento'])}/{md(n['mantenimiento'])}/{md(n['minimo'])}" for n in p["niveles"]))
                    for p in fl["patrones"]))
        L += ["", "### Niveles de presión por campaña (hojas Flow)", ""]
        L += table(["Campaña", "Hoja", "Lanzamiento", "Mantenimiento", "Mínimo"],
                   ((n["campana"], n["hoja"], n["lanzamiento"], n["mantenimiento"], n["minimo"]) for n in fl["niveles_por_campana"]))
        L += ["", "### Registros con TRPs o patrón", ""]
        L += table(["Hoja", "Fila", "Campaña", "Medio", "Patrón", "TRPs", "Semanas", "Prom. sem.", "Suma serie"],
                   ((f["hoja"], f["fila"], f["campana"], f["medio"], f["patron"], f["trps_totales"], f["semanas_al_aire"],
                     f["promedio_semanal"], f["suma_serie"]) for f in fl["registros"]), 60)
        v = fl["validaciones"]
        L += ["", f"### Validaciones ({len(v['trps_vs_serie_semanal'])} + {len(v['promedio_vs_trps_entre_semanas'])} sin cuadrar)", ""]
        L += table(["Dónde", "TRPs totales", "Suma serie semanal"],
                   ((x["donde"], x["trps_totales"], x["suma_serie"]) for x in v["trps_vs_serie_semanal"]))
        L += [""] + table(["Dónde", "Promedio semanal", "TRPs / semanas"],
                          ((x["donde"], x["promedio_semanal"], x["trps_entre_semanas"]) for x in v["promedio_vs_trps_entre_semanas"]))
        L += ["", "## 5. Ajustes de presupuesto", ""]
        L += table(["Celda", "Tipo", "Concepto", "Monto", "Fecha en texto"],
                   ((f"{a['hoja']}!{a['celda']}", a["tipo"], a["concepto"], a["monto"], a["fecha_texto"])
                    for a in rep["ajustes_presupuesto"] or []))
        n_hard = "6"
    else:
        L += ["## 7. Reglas implícitas en fórmulas (constantes recurrentes)", ""]
        n_hard = "7b"
    consts_md = table(["Valor", "Veces", "Operación", "Lectura tentativa", "Hojas", "Ejemplo"],
                      ((c["valor"], c["frecuencia"], ", ".join(f"{k}×{v}" for k, v in c["operaciones"].items()),
                        c["interpretacion"], ", ".join(c["hojas"]), f"{c['celdas'][0]}: `{c['ejemplo']}`")
                       for c in rep["constantes"]), 60)
    text_md = table(["Texto", "%", "Veces", "Celdas"],
                    ((t["texto"][:90], ", ".join(t["porcentajes"]), t["frecuencia"], ", ".join(t["celdas"][:3]))
                     for t in rep["reglas_en_texto"]), 40)
    if s["formato"] == "flow_anual":
        L += ["", "## 6. Fórmulas con números escritos a mano (deuda técnica)", ""]
    else:
        L += consts_md + ["", "### Porcentajes declarados en textos y encabezados", ""] + text_md
        L += ["", f"### {n_hard}. Fórmulas con números escritos a mano (deuda técnica)", ""]
    h = rep["formulas_hardcode"]
    L += [f"{h['total']} fórmulas.", ""]
    L += table(["Hoja", "Celda", "Números", "Fórmula"],
               ((x["hoja"], x["celda"], ", ".join(x["numeros"]), f"`{x['formula'][:100]}`") for x in h["formulas"]), 100)
    if s["formato"] == "flow_anual":
        L += ["", "### Constantes recurrentes", ""] + consts_md + ["", "### Porcentajes declarados en textos", ""] + text_md
        return "\n".join(L) + "\n"

    cat = rep["catalogos"]
    L += ["", "## 8. Catálogos", ""]
    su = cat.get("sumario")
    if su:
        L += [f"### Proveedores ({su['hoja']})", ""]
        L += table(["Medio", "Proveedor", "Razón social"], ((p["medio"], p["proveedor"], p["razon_social"]) for p in su["proveedores"]), 80)
        L += ["", "Tipos de compra: " + ", ".join(f"{k} ({v})" for k, v in su["tipos_de_compra"].items()), ""]
    for label, title in (("listas", "Listas"), ("radio", "Radio"), ("planta_ooh", "OOH (planta)"), ("digital", "Digital")):
        c = cat.get(label)
        if not c:
            continue
        L += [f"### {title} ({c['hoja']})", ""]
        for f, vals in c.items():
            if f != "hoja":
                shown = list(vals.items())[:25]
                L.append(f"- **{f}** ({len(vals)}): " + (", ".join(f"{md(k)} ({v})" for k, v in shown) or "—")
                         + (" …" if len(vals) > 25 else ""))
        L.append("")
    b = rep["bitacora"]
    L += ["## 9. Bitácora", "", f"- Hoja: {b['hoja']} · fila de registro: {md(b['fila_registro'])}",
          f"- Campos de encabezado: {', '.join(x['campo'] + (' ✔' if x['lleno'] else ' (vacío)') for x in b['campos_encabezado'])}",
          f"- Campos por entrada: {', '.join(b['campos_registro']) or '—'}",
          f"- Campos esperados: {', '.join(k + (' ✔' if v else ' ✘') for k, v in b.get('campos_esperados', {}).items())}",
          f"- Entradas llenas: **{b['entradas_llenas']}**", "", "## 10. Cadena de dependencias (hoja Flow)", ""]
    L += table(["Bloque", "Filas", "Hojas que lo alimentan (referencias)", "Externas"],
               ((d["bloque"], d["filas"], ", ".join(f"{k} ({v})" for k, v in d["fuentes"].items()) or "solo cálculos internos",
                 ", ".join(d["externas"]) or "—") for d in rep["dependencias"]))
    return "\n".join(L) + "\n"


def scrubber(terms):
    pats = [re.compile(r"(?<!\w)" + re.escape(t) + r"(?!\w)", re.I) for t in sorted(terms, key=len, reverse=True)]

    def scrub(s):
        s = str(s)
        for p in pats:
            s = p.sub("‹cliente›", s)
        return re.sub(r"\d[\d,]{4,}(\.\d+)?", "‹cifra›", s)  # montos: 5+ dígitos
    return scrub


def build_catalog(out):
    """catalog.md se regenera con TODOS los *.rules.json de la carpeta: reprocesar un archivo solo cambia su aporte."""
    reps = []
    for f in sorted(out.glob("*.rules.json")):
        try:
            r = json.loads(f.read_text(encoding="utf-8"))
        except Exception as e:
            log.warning("No se pudo leer %s: %s", f.name, e)
            continue
        if r["source"]["formato"] != "unknown":
            reps.append(r)
    scrub = scrubber({t for r in reps for t in r.get("terminos_cliente", [])})

    medios, patterns, consts, texts = {}, {}, {}, Counter()
    for r in reps:
        for m in r.get("medios") or []:
            d = medios.setdefault(m["medio"], {"variantes": Counter(), "archivos": 0})
            d["variantes"].update(m["variantes"])
            d["archivos"] += 1
        for p in (r.get("flighting") or {}).get("patrones", []):
            d = patterns.setdefault(p["patron"], {"variantes": Counter(), "campanas": 0, "filas": 0, "medios": Counter(), "niveles": []})
            d["variantes"].update(p["variantes"])
            d["campanas"] += p["campanas"]
            d["filas"] += p["filas"]
            d["medios"].update(p["medios"])
            d["niveles"] += p["niveles"]
        for c in r.get("constantes", []):
            if (abs(c["valor"]) >= 10_000 and c["valor"] not in (1e4, 1e5, 1e6, 1e7)) or \
                    (set(c["operaciones"]) <= {"+", "-"} and abs(c["valor"]) >= 100):
                continue  # ponytail: montos y ajustes sumados a mano, no reglas; quedan en la deuda técnica del JSON
            d = consts.setdefault(f"{c['valor']:g}", {"valor": c["valor"], "frecuencia": 0, "operaciones": Counter(),
                                                      "hojas": 0, "archivos": 0, "interpretacion": c["interpretacion"]})
            d["frecuencia"] += c["frecuencia"]
            d["operaciones"].update(c["operaciones"])
            d["hojas"] += len(c["hojas"])
            d["archivos"] += 1
        for t in r.get("reglas_en_texto", []):
            texts[scrub(t["texto"])[:100]] += t["frecuencia"]

    rng = lambda xs: "—" if not xs else f"{min(xs):g}" if min(xs) == max(xs) else f"{min(xs):g} a {max(xs):g}"
    L = ["# Reglas de planeación por confirmar", "",
         f"Documento generado a partir de {len(reps)} flowchart(s) el {dt.date.today().isoformat()}. "
         "Cada punto es una pregunta: marca la casilla si es correcto o escribe la corrección debajo. "
         "No incluye inversiones ni nombres de cliente; los conteos indican cuántas veces aparece cada regla.", ""]

    L += ["## 1. ¿Estos nombres son el mismo medio?", "",
          "El generador de flows necesita un catálogo único. Encontramos estos medios y la forma en que están escritos:", ""]
    for i, (k, d) in enumerate(sorted(medios.items(), key=lambda kv: -sum(kv[1]["variantes"].values())), 1):
        vs = ", ".join(f"«{scrub(v)}» ({n})" for v, n in d["variantes"].most_common())
        q = "¿Son el mismo medio?" if len(d["variantes"]) > 1 else "¿Es un medio vigente del catálogo?"
        L += [f"{i}. **{scrub(k)}** — {vs}. {q}", "   - [ ] Sí  - [ ] No → corrección: ____", ""]
    if not medios:
        L += ["Sin medios detectados.", ""]

    L += ["## 2. ¿Así se leen los patrones de flighting?", "",
          "Supuesto a confirmar: **MMTO y MANT son lo mismo (Mantenimiento)**, LANZ = Lanzamiento, MIN = Mínimo, "
          "y el separador «|» o «/» indica cambio de etapa en el tiempo.", "   - [ ] Correcto  - [ ] Corrección: ____", ""]
    for i, (k, d) in enumerate(sorted(patterns.items(), key=lambda kv: -kv[1]["campanas"]), 1):
        n = d["niveles"]
        lv = (f" Niveles de presión (TRPs semanales) en esas campañas: lanzamiento {rng([x['lanzamiento'] for x in n if x['lanzamiento'] is not None])}, "
              f"mantenimiento {rng([x['mantenimiento'] for x in n if x['mantenimiento'] is not None])}, "
              f"mínimo {rng([x['minimo'] for x in n if x['minimo'] is not None])}.") if n else " Sin niveles de presión registrados."
        med = ", ".join(f"{scrub(m)} ({c})" for m, c in d["medios"].most_common())
        L += [f"{i}. **{k}** — {d['campanas']} campaña(s), {d['filas']} fila(s); medios: {med}. "
              f"Escrito como: {', '.join('«' + scrub(v) + '»' for v in d['variantes'])}.{lv}",
              "   - ¿En qué orden y con qué duración se aplica cada etapa? ____", ""]
    if not patterns:
        L += ["Sin patrones detectados.", ""]

    L += ["## 3. ¿Qué significan estos factores fijos en las fórmulas?", "",
          "Son números escritos a mano que se repiten dentro de fórmulas. La lectura es una hipótesis nuestra.", "",
          "| # | Factor | Veces | Hojas | Archivos | Lectura tentativa | ¿Correcto? / regla real |", "|---|---|---|---|---|---|---|"]
    for i, d in enumerate(sorted(consts.values(), key=lambda d: -d["frecuencia"]), 1):
        ops = ", ".join(f"{o}{fv(d['valor'])}" for o in d["operaciones"])
        L.append(f"| {i} | {ops} | {d['frecuencia']} | {d['hojas']} | {d['archivos']} | {md(d['interpretacion'])} | ☐ ____ |")
    if not consts:
        L.append("| — | Sin constantes recurrentes | | | | | |")

    L += ["", "## 4. ¿Siguen vigentes estos porcentajes escritos en encabezados y notas?", ""]
    for i, (t, n) in enumerate(texts.most_common(30), 1):
        L.append(f"{i}. «{md(t)}» — {n} vez/veces. ☐ Vigente ☐ Cambió a ____")
    if not texts:
        L.append("Sin porcentajes declarados.")
    (out / "catalog.md").write_text("\n".join(L) + "\n", encoding="utf-8")


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("path", nargs="?", default=DATA_DIR, help=f"archivo o carpeta (default: {DATA_DIR})")
    ap.add_argument("--out", default=OUT_DIR, help=f"carpeta de salida (default: {OUT_DIR})")
    args = ap.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", datefmt="%H:%M:%S")

    root, out = Path(args.path).expanduser().resolve(), Path(args.out).expanduser().resolve()
    if not root.exists():
        log.error("No existe: %s", root)
        return 2
    files = [root] if root.is_file() else sorted(
        p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in EXTS and not p.name.startswith("~$"))
    if not files:
        log.warning("No hay archivos .xlsx/.xlsm en %s", root)
        return 0
    out.mkdir(parents=True, exist_ok=True)

    failed = 0
    for i, f in enumerate(files, 1):
        k = f.name if root.is_file() else "__".join(f.relative_to(root).parts)
        log.info("[%d/%d] %s", i, len(files), f)
        try:
            process(f, k, out)
        except Exception:
            failed += 1
            log.exception("  no se pudo procesar %s", f)
    build_catalog(out)
    log.info("Listo: %d archivos, %d con error. Salida en %s", len(files), failed, out)
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
