"""Self-check: python test_extract_flow.py"""
import datetime as dt
import json
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter as L

import extract_flow as p

WEEKS = [dt.datetime(2025, 1, 6) + dt.timedelta(days=7 * i) for i in range(52)]
FIRST_WEEK = 7  # columna G


def annual(path, skew=0):
    """Flow anual mínimo: calendario de 52 semanas en la fila 7 (no la 47), dos categorías."""
    wb = openpyxl.Workbook()
    fl = wb.active
    fl.title = "FLOW"
    fl["A2"], fl["C2"] = "PPTO MKT", "=F8+33000000"
    for c, q in ((0, "1Q"), (13, "2Q"), (26, "3Q"), (39, "4Q")):
        fl.cell(5, FIRST_WEEK + c, q)
    for c, m in ((0, "ENERO"), (4, "FEBRERO"), (8, "MARZO")):
        fl.cell(6, FIRST_WEEK + c, m)
    tail = FIRST_WEEK + 52
    for c, h in enumerate(["CATEGORIA", "CAMPAÑAS", "PRIORIDAD", "TARGET", "MEDIOS", "INVERSIÓN"], 1):
        fl.cell(7, c, h)
    for i, d in enumerate(WEEKS):
        fl.cell(7, FIRST_WEEK + i, d)
    for c, h in enumerate(["TRP´S TOTALES", "WKS PROM SEM", "PESO CAMPAÑA", "SEMANAS AL AIRE"]):
        fl.cell(7, tail + c, h)

    def block(r, cat, camp, target, medios, patterns, weeks_on_air):
        if cat:
            fl.cell(r, 1, cat)
        fl.cell(r, 2, camp)
        fl.cell(r, 4, target)
        for k, (m, pat) in enumerate(zip(medios, patterns)):
            fl.cell(r + k, 5, m)
            fl.cell(r + k, 6, 1_000_000)
            if pat:
                fl.cell(r + k, tail + 2, pat)
            fl.cell(r + k, tail + 3, weeks_on_air)
        for i in range(weeks_on_air):  # TRPs semanales en la fila de TV
            fl.cell(r, FIRST_WEEK + i, 100)
        fl.cell(r, tail, 100 * weeks_on_air + skew)
        fl.cell(r, tail + 1, 100)
        fl.merge_cells(start_row=r, end_row=r + len(medios) - 1, start_column=2, end_column=2)

    block(8, "CAT_A", "CAMP UNO", "P 25-45", ["TV Abierta", "Digital", "OOH"],
          ["MANT | LANZ | MIN", "MMTO / LANZ / MIN", None], 10)
    block(11, None, "CAMP DOS", "P 18-35", ["TV Abierta", "Digital", "OOH"], ["MANT  / MIN", None, None], 8)
    fl.merge_cells("A8:A13")
    fl["A14"] = "TOTAL CAT_A"
    block(15, "CAT_B", "CAMP TRES", "M 25-45", ["Tv Abierta", "Digital", "OOH"], ["LANZ|MMTO", None, None], 5)
    fl.merge_cells("A15:A17")
    fl["A18"] = "TOTAL CAT_B"
    fl["A19"] = "GRAN TOTAL"

    cat = wb.create_sheet("Flow CAT 1")
    for c, h in enumerate(["CATEGORIA", "CAMPAÑAS", "MEDIOS", "INVERSIÓN x Medio", "INVERSIÓN"], 2):
        cat.cell(7, c, h)
    for i, d in enumerate(WEEKS):
        cat.cell(7, FIRST_WEEK + i, d)
    for c, h in enumerate(["TRPs TOTALES", "WEEKLYS", "WOA", None, "Lanzamiento", "Mantenimiento", "Mínimo"]):
        if h:
            cat.cell(7, tail + c, h)
    for r, (camp, lv) in enumerate([("CAMP UNO", (190, 150, 100)), ("CAMP TRES", (180, 140, 90))]):
        row = 8 + r * 2
        cat.cell(row, 2, "CAT_A" if r == 0 else "CAT_B")
        cat.cell(row, 3, camp)
        cat.cell(row, 4, "TV Abierta")
        cat.cell(row + 1, 4, "Inversión por periodo")
        for k, v in enumerate(lv):
            cat.cell(row, tail + 4 + k, v)
    wb.create_sheet("Resumen Escenario")
    eb = wb.create_sheet("EXTRA BUDGET-LINE")
    for row in [("Media", "Period", "Investment"), ("PPTO ADICIONAL CAMP UNO (DIG) 10JUL", None, 250000),
                ("RECORTE General16 oct", None, -500000), (None, "TOTAL", -250000)]:
        eb.append((None,) + row)
    wb.save(path)


def campaign(path):
    wb = openpyxl.Workbook()
    b = wb.active
    b.title = "Bitacora"
    for row in [("Cliente: ", "ClienteSecreto"), ("Campaña:", "CampX"), (), (), (), ("Versión ", 1), (),
                ("Fecha ", "Cambio ", "Persona que Solicito ", "Persona que envio ", "Inversión del Flow"),
                ("2026-04-01", "Alta", "Ana", "Luis", 100), ("2026-04-08", "Ajuste", "Ana", "Luis", 120), (None, " ")]:
        b.append(row)
    wb.create_sheet("Kpis")["A1"] = "KPIs"
    fl = wb.create_sheet("Flow")
    fl["B1"], fl["C1"] = "Cliente: ", "ClienteSecreto"
    fl["B8"] = "Media"
    fl["B10"], fl["C12"], fl["C13"] = "TV ABIERTA ", "=TVA!N11*2", "=TVA!N12*2"
    fl["J2"], fl["K2"], fl["Q2"], fl["R2"] = 10, "=+J2*66%", 20, "=+Q2*50%"
    fl["B20"], fl["C21"] = "RADIO ", "=Radio!F12+Radio!F13"
    fl["B30"] = "DIGITAL "
    for i, f in enumerate(["/7", "/7", "/7", "/4", "/4", "/3", "/3"]):
        fl.cell(31, 5 + i, f"='DIGITAL '!M{5 + i}{f}")
    tva = wb.create_sheet("TVA")
    tva.append(["Canal", "Tarifa 2B", "Tarifa 2B DESC 20% POR VN", "Inflación 5%"])
    for r in range(2, 5):
        tva.append(["2 NAC", 1000, f"=B{r}*0.8", 7])
    radio = wb.create_sheet("Radio")
    radio.append(["PLAZA", "SIGLAS", "Proveedor", "Formato_HM", "Tipo", "Tarifa"])
    for plaza, sig, base in [("DF", "XEX-FM", 128), ("GDL", "XEW-FM", 90), ("MTY", "XHM-FM", 60)]:
        radio.append([plaza, sig, "Grupo Radio", 'SPOT 20" (3-20)', "Mandatorio", f"={base}*1.04"])
    su = wb.create_sheet("Sumario")
    su.append(["Cliente: ", "ClienteSecreto"])
    su.append([])
    su.append(["MEDIO ", "PROVEEDOR", "RAZÓN SOCIAL ", "TIPO PAGO"])
    for row in [("TV ABIERTA", "PRV 1", "Televisora SA", "PAP"), (None, "PRV 2", "Otra SA", "FD"),
                ("RADIO", "PRV 3", "ProveedorSecreto SA", "DIRECTO"), ("DIGITAL", "PRV 4", "Digital SA", "BONUS")]:
        su.append(row)
    li = wb.create_sheet("Listas")
    li.append(["Tipos Facturacion", "Formato", "Spots"])
    li.append(["FD (1)", 'SPOT 20" (3-20)', 8])
    li.append(["PAP (3)", 'SPOT 30" (3-30)', 9])
    pl = wb.create_sheet("PLANTA.")
    pl.append(["Planta", "Tipo de Compra", "Tipo", "Formato", "Periodo"])
    pl.append(["Fija", "Planta", "Fijo", "Cartelera", "Mensual"])
    pl.append(["Fija", "Planta", "Digital", "Pantalla", "Otro"])
    dg = wb.create_sheet("DIGITAL ")
    dg.append(["Tipo de Pauta", "Medio", "Plataforma", "Placement", "Tipo de Compra", "KPI"])
    dg.append(["CTV REGULAR", "Medio 1", "DIRECT BUY", "CTV", "CPM", "Impresiones"])
    dg.append(["PODCAST", "Medio 2", "Medio 9", "All Devices", "FIXED", "N/A"])
    wb.save(path)


def load(out, name):
    return json.loads((out / f"{name}.rules.json").read_text())


def test():
    tmp = Path(tempfile.mkdtemp())
    data, out = tmp / "data", tmp / "out"
    data.mkdir()
    annual(data / "anual.xlsx")
    annual(data / "descuadre.xlsx", skew=37)
    campaign(data / "campana.xlsx")
    wb = openpyxl.Workbook()
    wb.active["A1"] = "nada"
    wb.save(data / "otro.xlsx")
    assert p.main([str(data), "--out", str(out)]) == 0

    r = load(out, "anual.xlsx")
    assert r["source"]["formato"] == "flow_anual"
    cal = r["calendario"]
    assert (cal["fila"], cal["fila_meses"], cal["fila_trimestres"], len(cal["semanas"])) == (7, 6, 5, 52)
    w = cal["semanas"]
    assert (w[0]["fecha"], w[0]["columna"], w[0]["mes"], w[13]["trimestre"], w[51]["columna"]) == \
        ("2025-01-06", "G", "ENERO", "2Q", L(FIRST_WEEK + 51))
    tv = next(m for m in r["medios"] if m["medio"] == "TV ABIERTA")
    assert tv["variantes"] == {"TV Abierta": 4, "Tv Abierta": 1}, tv  # FLOW (2 + 1) + Flow CAT 1 (2)
    combos = {(h["categoria"], h["campana"], h["target"]) for h in r["jerarquia"]["combinaciones"] if "FLOW" in h["hojas"]}
    assert combos == {("CAT_A", "CAMP UNO", "P 25-45"), ("CAT_A", "CAMP DOS", "P 18-35"), ("CAT_B", "CAMP TRES", "M 25-45")}
    pats = {x["patron"]: x for x in r["flighting"]["patrones"]}
    assert set(pats) == {"MANT | LANZ | MIN", "MANT | MIN", "LANZ | MANT"}, set(pats)
    assert pats["MANT | LANZ | MIN"]["variantes"] == {"MANT | LANZ | MIN": 1, "MMTO / LANZ / MIN": 1}
    assert pats["MANT | LANZ | MIN"]["campanas"] == 1 and pats["MANT | LANZ | MIN"]["filas"] == 2
    assert pats["MANT | LANZ | MIN"]["niveles"] == [{"lanzamiento": 190, "mantenimiento": 150, "minimo": 100}]
    assert pats["LANZ | MANT"]["niveles"] == [{"lanzamiento": 180, "mantenimiento": 140, "minimo": 90}]
    v = r["flighting"]["validaciones"]
    assert v == {"trps_vs_serie_semanal": [], "promedio_vs_trps_entre_semanas": []}, v
    adj = [(a["tipo"], a["monto"], a["fecha_texto"]) for a in r["ajustes_presupuesto"]]
    assert adj == [("adicional", 250000, "10-JUL"), ("recorte", -500000, "16-OCT")], adj
    assert [h["celda"] for h in r["formulas_hardcode"]["formulas"]] == ["C2"]

    bad = load(out, "descuadre.xlsx")["flighting"]["validaciones"]
    assert len(bad["trps_vs_serie_semanal"]) == 3 and len(bad["promedio_vs_trps_entre_semanas"]) == 3

    r = load(out, "campana.xlsx")
    assert r["source"]["formato"] == "flow_campana"
    consts = {(c["valor"], op): c["frecuencia"] for c in r["constantes"] for op in c["operaciones"]}
    assert consts[(7, "/")] == 3 and consts[(4, "/")] == 2 and consts[(3, "/")] == 2
    assert consts[(0.8, "*")] == 3 and consts[(1.04, "*")] == 3 and consts[(0.66, "*")] == 1 and consts[(0.5, "*")] == 1
    assert (128, "*") not in consts  # tarifa capturada, no factor
    texts = {t["texto"] for t in r["reglas_en_texto"]}
    assert {"Tarifa 2B DESC 20% POR VN", "Inflación 5%"} <= texts
    cat = r["catalogos"]
    assert cat["sumario"]["tipos_de_compra"] == {"PAP": 1, "FD": 1, "DIRECTO": 1, "BONUS": 1}
    assert [x["medio"] for x in cat["sumario"]["proveedores"]] == ["TV ABIERTA", "TV ABIERTA", "RADIO", "DIGITAL"]
    assert cat["radio"]["plazas"] == {"DF": 1, "GDL": 1, "MTY": 1} and len(cat["radio"]["estaciones"]) == 3
    assert cat["planta_ooh"]["formatos"] == {"Cartelera": 1, "Pantalla": 1} and cat["planta_ooh"]["tipos"] == {"Fijo": 1, "Digital": 1}
    assert cat["digital"]["plataformas"] == {"DIRECT BUY": 1, "Medio 9": 1} and cat["digital"]["tipos_de_pauta"] == {"CTV REGULAR": 1, "PODCAST": 1}
    assert cat["listas"]["SPOTS"] == {"8": 1, "9": 1}
    b = r["bitacora"]
    assert b["entradas_llenas"] == 2 and all(b["campos_esperados"].values()) and b["version_llena"]
    deps = {d["bloque"]: d["fuentes"] for d in r["dependencias"]}
    assert deps == {"TV ABIERTA": {"TVA": 2}, "RADIO": {"Radio": 2}, "DIGITAL": {"DIGITAL ": 7}}, deps

    assert load(out, "otro.xlsx")["source"]["formato"] == "unknown"

    catalog = (out / "catalog.md").read_text()
    for leak in ("ClienteSecreto", "CAMP UNO", "CAT_A", "ProveedorSecreto", "250000", "250,000", "33000000", "33,000,000"):
        assert leak not in catalog, leak
    for must in ("MANT | LANZ | MIN", "«Tv Abierta»", "| /7 |", "*1.04", "Inflación 5%", "lanzamiento 190"):
        assert must in catalog, must
    print("ok")


if __name__ == "__main__":
    test()
