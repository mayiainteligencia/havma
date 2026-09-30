"""Self-check: python test_profile_xlsx.py"""
import json
import logging
import tempfile
from pathlib import Path

import openpyxl

import profile_xlsx as p


def test():
    tmp = Path(tempfile.mkdtemp())
    data, out = tmp / "data", tmp / "reports"
    (data / "sub").mkdir(parents=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws.append(["Plan de medios"])
    ws.append(["Medio", "Inversion", "Fecha"])
    ws.append(["TV", 100, None])
    ws.append(["Radio", 50, None])
    ws["D3"] = "#REF!"
    ws.merge_cells("A1:C1")
    ws.row_dimensions[4].hidden = True
    ws.column_dimensions["E"].hidden = True
    out_ws = wb.create_sheet("Salida")
    out_ws["A1"] = "=Inputs!B3+33000000"
    out_ws["A2"] = "=ROUND('Inputs'!B4*1.16,2)"
    out_ws["A3"] = "=[1]Hoja!A1"
    out_ws["A4"] = "=SUM(A1:A2)"
    kv = wb.create_sheet("Parametros")
    for row in [("Cliente", "Spectrum"), ("Target", "AB 25-44"), ("Universo", 17420700), (None, None),
                ("Version", 2), ("Fecha", "2026-01-15")]:
        kv.append((None,) + row)  # columnas B:C, como la hoja Exercise real
    wb.save(data / "sub" / "plan.xlsx")
    (data / "~$plan.xlsx").write_text("lock")
    (data / "c.csv").write_text("medio;monto;semanas\nTV;1,000.5;4\nRadio;20;4\n", encoding="utf-8")

    assert p.main([str(data), "--out", str(out)]) == 0
    rep = json.loads((out / "sub__plan.xlsx.profile.json").read_text())
    inp, sal, par = rep["sheets"]
    assert rep["cached_values"] == "missing"
    assert inp["layout"] == "table" and par["layout"] == "key_value" and par["column_profile"] == []
    assert [(x["row"], x["key"], x["value"]) for x in par["key_values"]] == [
        (1, "Cliente", "Spectrum"), (2, "Target", "AB 25-44"), (3, "Universo", 17420700),
        (5, "Version", 2), (6, "Fecha", "2026-01-15")]
    assert "| Parametros | visible | key_value |" in (out / "sub__plan.xlsx.profile.md").read_text()
    assert inp["header_row"] == 2 and inp["column_profile"][1]["dominant_type"] == "number"
    assert inp["errors"]["by_type"] == {"#REF!": 1}
    assert inp["merged_cells"]["total"] == 1 and inp["hidden_rows"]["ranges"] == ["4"]
    assert inp["hidden_columns"]["ranges"] == ["E"]
    assert sal["formulas"] == 4 and sal["hardcoded_formulas"]["total"] == 2
    assert [c["numbers"] for c in sal["hardcoded_formulas"]["cells"]] == [["33000000"], ["1.16"]]
    assert sal["references"]["cross_sheet"] == 2 and sal["references"]["external"] == 1
    assert sal["references"]["internal"] == 1
    assert {"from": "Inputs", "to": "Salida", "refs": 2} in rep["lineage"]["edges"]
    assert rep["lineage"]["input_sheets"] == ["Inputs", "Parametros"]
    assert rep["lineage"]["output_sheets"] == ["Salida", "Parametros"]
    csv_rep = json.loads((out / "c.csv.profile.json").read_text())
    assert csv_rep["sheets"][0]["column_profile"][1]["max"] == 1000.5

    readme = out / "README.md"
    readme.write_text(readme.read_text() + "\nNota manual del equipo.\n")
    assert p.main([str(data / "sub" / "plan.xlsx"), "--out", str(out)]) == 0  # clave distinta: plan.xlsx
    assert p.main([str(data), "--out", str(out)]) == 0
    text = readme.read_text()
    assert text.count("<!-- profile:sub__plan.xlsx -->") == 1 and "Nota manual del equipo." in text
    assert "~$" not in text

    # Si openpyxl cambia su atributo privado, el perfil sigue y avisa una sola vez.
    warnings = []
    handler = logging.Handler()
    handler.emit = lambda rec: warnings.append(rec) if rec.levelno == logging.WARNING else None
    p.log.addHandler(handler)
    real_layout, p._layout_warned = p.sheet_layout, False
    p.sheet_layout = lambda path, part: (_ for _ in ()).throw(AttributeError("_worksheet_path"))
    try:
        assert p.main([str(data), "--out", str(out)]) == 0
    finally:
        p.sheet_layout = real_layout
        p.log.removeHandler(handler)
    rep = json.loads((out / "sub__plan.xlsx.profile.json").read_text())
    assert all(s[f] == "no disponible" for s in rep["sheets"] for f in ("merged_cells", "hidden_rows", "hidden_columns"))
    assert rep["sheets"][0]["errors"]["total"] == 1 and rep["sheets"][2]["layout"] == "key_value"
    assert len([w for w in warnings if "no disponibles" in w.getMessage()]) == 1
    print("ok")


if __name__ == "__main__":
    test()
