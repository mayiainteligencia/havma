#!/usr/bin/env python3
"""Data profiling y linaje de fórmulas para libros Excel de planeación de medios.

Uso:
    python profile_xlsx.py [ruta_archivo_o_carpeta] [--out ./reports]

Sin ruta procesa DATA_DIR (o $HAVAS_DATA_DIR). Dependencias: openpyxl + stdlib.
"""
import argparse
import csv
import datetime as dt
import json
import logging
import os
import re
import sys
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter, defaultdict
from itertools import zip_longest
from pathlib import Path
from urllib.parse import quote

import openpyxl
from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter

# Único lugar a cambiar al mover el proyecto a la nube privada (o exportar HAVAS_DATA_DIR).
DATA_DIR = os.environ.get("HAVAS_DATA_DIR", "/Users/martinfranciscocortesbuendia/Desktop/EdgeNET/havasData/dataEx")
REPORTS_DIR = Path(__file__).resolve().parent / "reports"

EXTS = {".xlsx", ".xlsm", ".csv"}
ERROR_CODES = ("#REF!", "#VALUE!", "#N/A", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!")
MAX_ERROR_CELLS = 50
MAX_DETAIL = 200     # ponytail: tope de ejemplos por hoja (hardcodes/externas); los conteos siempre son completos
HEADER_SCAN = 50     # ponytail: el encabezado se busca solo en las primeras 50 filas con contenido
UNIQUE_CAP = 10_000  # por arriba de esto se reporta ">10000" para no agotar memoria
MAX_EXAMPLES = 10
KV_MIN_UNIQUE = 0.7  # hoja clave-valor: 2a columna con ≥70% valores únicos (1.0 = sin ningún repetido)

log = logging.getLogger("profile_xlsx")
NA = "no disponible"
_layout_warned = False

CELL_RE = re.compile(r"^[A-Z]{1,3}\d+(:[A-Z]{1,3}\d+)?$|^[A-Z]{1,3}:[A-Z]{1,3}$|^\d+:\d+$", re.I)
EXT_RE = re.compile(r"\[([^\]]+)\]")
NUM_RE = re.compile(r"^-?\$?\d[\d,]*(\.\d+)?$|^-?\$?\.\d+$")
ARITH = set("+-*/^")


# ---------------------------------------------------------------- utilidades

def cell_type(v):
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return "bool"
    if isinstance(v, (int, float)):
        return "number"
    if isinstance(v, (dt.datetime, dt.date)):
        return "date"
    if isinstance(v, dt.time):
        return "time"
    if isinstance(v, dt.timedelta):
        return "duration"
    if isinstance(v, str):
        return "error" if v in ERROR_CODES else "text"
    return "other"


def ranges(nums):
    """[1,2,3,7] -> ['1-3', '7']"""
    out, nums = [], sorted(set(nums))
    for n in nums:
        if out and out[-1][1] == n - 1:
            out[-1][1] = n
        else:
            out.append([n, n])
    return [f"{a}-{b}" if a != b else str(a) for a, b in out]


def md(s):
    return str(s).replace("|", "\\|").replace("\n", " ")


# ---------------------------------------------------------------- fórmulas

def parse_formula(text):
    """Devuelve (operandos de rango, números escritos a mano en operaciones aritméticas)."""
    toks = [t for t in Tokenizer(text).items if t.type != "WHITE-SPACE"]
    only_numbers = all(t.type == "OPERATOR-PREFIX" or t.subtype == "NUMBER" for t in toks)
    refs, nums = [], []
    for i, t in enumerate(toks):
        if t.type != "OPERAND":
            continue
        if t.subtype == "RANGE":
            refs.append(t.value)
        elif t.subtype == "NUMBER":
            j = i - 1
            while j >= 0 and toks[j].type == "OPERATOR-PREFIX":
                j -= 1
            neighbours = (toks[j] if j >= 0 else None, toks[i + 1] if i + 1 < len(toks) else None)
            # ponytail: solo números junto a + - * / ^ (=A1+33000000, =A1*1.16); args como ROUND(x,2) no cuentan
            if only_numbers or any(n is not None and n.type == "OPERATOR-INFIX" and n.value in ARITH for n in neighbours):
                nums.append(t.value)
    return refs, nums


def classify_ref(tok, own, sheets_lc, ext_links, names):
    """-> (tipo, destinos) con tipo en internal | cross_sheet | external | name."""
    if "!" in tok:
        sheet, _ = tok.rsplit("!", 1)
        if sheet.startswith("'") and sheet.endswith("'"):
            sheet = sheet[1:-1].replace("''", "'")
        m = EXT_RE.search(sheet)
        if m:
            book = ext_links.get(m.group(1), m.group(1))
            return "external", [f"[ext] {book}"]
        sheet = sheets_lc.get(sheet.lower(), sheet)
        return ("internal", [own]) if sheet == own else ("cross_sheet", [sheet])
    if CELL_RE.match(tok.replace("$", "")):
        return "internal", [own]
    m = re.match(r"^\[(\d+)\]", tok)
    if m:
        return "external", [f"[ext] {ext_links.get(m.group(1), m.group(1))}"]
    name = tok.split("[")[0].upper()  # nombre definido o tabla estructurada Tabla1[Col]
    return "name", sorted(names.get(name, ()))


def defined_names(wb):
    names = defaultdict(set)
    scopes = [wb.defined_names] + [ws.defined_names for ws in wb.worksheets]
    for scope in scopes:
        for name, dn in scope.items():
            try:
                for sheet, _ in dn.destinations:
                    names[name.upper()].add(sheet)
            except Exception:  # nombres con fórmulas o referencias rotas
                pass
    return names


# ---------------------------------------------------------------- perfil de columnas

class Columns:
    def __init__(self):
        self.header_row, self.header, self.cols, self.rows = None, {}, {}, 0

    def set_header(self, r, vals):
        self.header_row = r
        self.header = {c: str(v).strip() for c, v in enumerate(vals, 1) if v not in (None, "")}

    def feed(self, vals):
        if all(v in (None, "") for v in vals):
            return
        self.rows += 1
        for c, v in enumerate(vals, 1):
            t = cell_type(v)
            if t is None:
                continue
            col = self.cols.get(c)
            if col is None:
                col = self.cols[c] = {"n": 0, "types": Counter(), "uniq": set(), "minmax": {}, "ex": []}
            col["n"] += 1
            col["types"][t] += 1
            if col["uniq"] is not None:
                col["uniq"].add(v)
                if len(col["uniq"]) > UNIQUE_CAP:
                    col["uniq"] = None
            if len(col["ex"]) < MAX_EXAMPLES and v not in col["ex"]:
                col["ex"].append(v)
            if t in ("number", "date", "text", "time", "duration"):
                mm = col["minmax"].get(t)
                if mm is None:
                    col["minmax"][t] = [v, v]
                else:
                    try:
                        if v < mm[0]:
                            mm[0] = v
                        if v > mm[1]:
                            mm[1] = v
                    except TypeError:  # date vs datetime
                        pass

    def result(self):
        out = []
        for c in sorted(set(self.cols) | set(self.header)):
            col = self.cols.get(c, {"n": 0, "types": Counter(), "uniq": set(), "minmax": {}, "ex": []})
            dom = col["types"].most_common(1)[0][0] if col["types"] else None
            mm = col["minmax"].get(dom, [None, None])
            out.append({
                "column": get_column_letter(c),
                "header": self.header.get(c, get_column_letter(c)),
                "dominant_type": dom,
                "type_counts": dict(col["types"]),
                "null_pct": round(100 * (self.rows - col["n"]) / self.rows, 2) if self.rows else None,
                "unique": len(col["uniq"]) if col["uniq"] is not None else f">{UNIQUE_CAP}",
                "min": mm[0],
                "max": mm[1],
                "examples": col["ex"],
            })
        return out


def find_header(buf):
    """Primera fila con mayoría de texto seguida de una fila con datos (algún valor no-texto)."""
    types = [[cell_type(v) for v in vals] for _, vals in buf]
    fallback = None
    for i, t in enumerate(types):
        filled = [x for x in t if x]
        if len(filled) < 2 or filled.count("text") * 2 <= len(filled):
            continue
        if fallback is None:
            fallback = i
        nxt = next((u for u in types[i + 1:] if any(u)), None)
        if nxt and any(x not in (None, "text") for x in nxt):
            return i
    return fallback


def key_value_cols(buf):
    """(col_clave, col_valor) si el buffer solo usa 2 columnas, la 1a mayoritariamente texto y la 2a casi sin repetidos."""
    used = sorted({c for _, vals in buf for c, v in enumerate(vals, 1) if v not in (None, "")})
    if len(used) != 2 or len(buf) < 2:
        return None
    k, v = used
    keys = [vals[k - 1] for _, vals in buf if len(vals) >= k and vals[k - 1] not in (None, "")]
    values = [repr(vals[v - 1]) for _, vals in buf if len(vals) >= v and vals[v - 1] not in (None, "")]
    if sum(cell_type(x) == "text" for x in keys) * 2 <= len(keys) or len(set(values)) < KV_MIN_UNIQUE * len(values):
        return None
    return k, v


# ---------------------------------------------------------------- hoja

def analyze_sheet(title, rows_f, rows_v, sheets_lc, ext_links, names):
    counts = Counter()
    err_count, err_cells = Counter(), []
    hardcoded, external = [], []
    kinds, sources, ext_books = Counter(), Counter(), Counter()
    bounds = [None, None, 0, 0]  # min_r, min_c, max_r, max_c
    cols, pending = Columns(), []
    kv, pairs = None, []

    def add_pair(r, vals):
        k, v = (vals[c - 1] if len(vals) >= c else None for c in kv)
        if k not in (None, "") or v not in (None, ""):
            pairs.append({"row": r, "key": k, "value": v})

    def settle():
        nonlocal kv
        kv = key_value_cols(pending)
        if kv:  # ponytail: después de la fila 50 solo se leen las 2 columnas detectadas
            for r, vals in pending:
                add_pair(r, vals)
            return
        i = find_header(pending)
        if i is not None:
            cols.set_header(*pending[i])
        for _, vals in pending[(i + 1 if i is not None else 0):]:
            cols.feed(vals)

    for r, (frow, vrow) in enumerate(zip_longest(rows_f, rows_v or (), fillvalue=()), start=1):
        vals = []
        for c, (f, v) in enumerate(zip_longest(frow, vrow), start=1):
            text = getattr(f, "text", f)  # ArrayFormula -> texto
            is_formula = isinstance(text, str) and text.startswith("=")
            if not is_formula and (f is None or f == ""):
                vals.append(None)
                continue
            val = v if is_formula else f
            vals.append(val)
            counts["non_empty"] += 1
            if bounds[0] is None:
                bounds[0] = r
            bounds[1] = c if bounds[1] is None else min(bounds[1], c)
            bounds[2], bounds[3] = r, max(bounds[3], c)
            coord = None
            if is_formula:
                counts["formulas"] += 1
                if v is None:
                    counts["formulas_without_cache"] += 1
                coord = f"{get_column_letter(c)}{r}"
                try:
                    refs, nums = parse_formula(text)
                except Exception:
                    counts["formula_parse_errors"] += 1
                    refs, nums = [], []
                if nums:
                    counts["hardcoded_formulas"] += 1
                    if len(hardcoded) < MAX_DETAIL:
                        hardcoded.append({"cell": coord, "formula": text, "numbers": nums})
                has_ext = False
                for ref in refs:
                    kind, targets = classify_ref(ref, title, sheets_lc, ext_links, names)
                    kinds[kind] += 1
                    if kind == "name" and not targets:
                        kinds["unresolved_name"] += 1
                    for t in targets:
                        if t != title:
                            sources[t] += 1
                        if kind == "external":
                            ext_books[t[6:]] += 1
                            has_ext = True
                if has_ext and len(external) < MAX_DETAIL:
                    external.append({"cell": coord, "formula": text})
            err = val if isinstance(val, str) and val in ERROR_CODES else None
            if err is None and is_formula and "#REF!" in text:
                err = "#REF!"
            if err:
                err_count[err] += 1
                if len(err_cells) < MAX_ERROR_CELLS:
                    err_cells.append({"cell": coord or f"{get_column_letter(c)}{r}", "error": err})

        if pending is None:
            add_pair(r, vals) if kv else cols.feed(vals)
        elif any(v not in (None, "") for v in vals):
            pending.append((r, vals))
            if len(pending) == HEADER_SCAN:
                settle()
                pending = None
    if pending is not None:
        settle()

    min_r, min_c, max_r, max_c = bounds
    n = counts["non_empty"]
    return {
        "name": title,
        "used_range": f"{get_column_letter(min_c)}{min_r}:{get_column_letter(max_c)}{max_r}" if n else None,
        "rows": max_r - min_r + 1 if n else 0,
        "columns": max_c - min_c + 1 if n else 0,
        "non_empty_cells": n,
        "formulas": counts["formulas"],
        "formula_pct": round(100 * counts["formulas"] / n, 2) if n else 0,
        "formulas_without_cache": counts["formulas_without_cache"],
        "formula_parse_errors": counts["formula_parse_errors"],
        "errors": {"total": sum(err_count.values()), "by_type": dict(err_count), "cells": err_cells},
        "hardcoded_formulas": {"total": counts["hardcoded_formulas"], "cells": hardcoded},
        "external_refs": {"total": kinds["external"], "books": dict(ext_books), "cells": external},
        "references": {
            "internal": kinds["internal"], "cross_sheet": kinds["cross_sheet"], "external": kinds["external"],
            "named": kinds["name"], "unresolved_names": kinds["unresolved_name"], "sources": dict(sources),
        },
        "layout": "key_value" if kv else "table",
        "key_values": pairs,
        "header_row": cols.header_row,
        "column_profile": cols.result(),
    }


def sheet_layout(path, part):
    """Celdas combinadas y filas/columnas ocultas leyendo el XML en streaming (read_only no las expone)."""
    merged, hrows, hcols, last = [], [], [], 0
    with zipfile.ZipFile(path) as z, z.open(part) as fh:
        for _, el in ET.iterparse(fh):
            tag = el.tag.rsplit("}", 1)[-1]
            if tag == "row":
                last = int(el.get("r") or last + 1)
                if el.get("hidden") in ("1", "true"):
                    hrows.append(last)
                el.clear()
            elif tag == "col" and el.get("hidden") in ("1", "true"):
                hcols.extend(range(int(el.get("min")), int(el.get("max")) + 1))
            elif tag == "mergeCell":
                merged.append(el.get("ref"))
    col_ranges = []
    for rg in ranges(hcols):
        a, _, b = rg.partition("-")
        col_ranges.append(get_column_letter(int(a)) + (f":{get_column_letter(int(b))}" if b else ""))
    return {
        "merged_cells": {"total": len(merged), "ranges": merged[:MAX_DETAIL]},
        "hidden_rows": {"total": len(hrows), "ranges": ranges(hrows)[:MAX_DETAIL]},
        "hidden_columns": {"total": len(hcols), "ranges": col_ranges[:MAX_DETAIL]},
    }


# ---------------------------------------------------------------- archivo

def lineage(sheets, sheet_names):
    edges = Counter()
    for s in sheets:
        for src, n in s["references"]["sources"].items():
            edges[(src, s["name"])] += n
    fed = {b for (a, b) in edges if a in sheet_names}
    consumed = {a for (a, b) in edges}
    names = [s["name"] for s in sheets]
    return {
        "edges": [{"from": a, "to": b, "refs": n} for (a, b), n in edges.most_common()],
        "input_sheets": [n for n in names if n not in fed],
        "output_sheets": [n for n in names if n not in consumed],
    }


def profile_xlsx(path):
    wb_f = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        wb_v = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        log.warning("  no se pudieron leer valores cacheados: %s", e)
        wb_v = None
    sheets, failures = [], []
    try:
        ext_links = {}
        for i, link in enumerate(getattr(wb_f, "_external_links", []), 1):
            target = getattr(getattr(link, "file_link", None), "Target", None) or f"[{i}]"
            ext_links[str(i)] = re.split(r"[\\/]", target)[-1]
        names = defined_names(wb_f)
        sheets_lc = {n.lower(): n for n in wb_f.sheetnames}
        for ws in wb_f.worksheets:
            log.info("  hoja '%s'", ws.title)
            try:
                rows_v = wb_v[ws.title].iter_rows(values_only=True) if wb_v else None
                s = analyze_sheet(ws.title, ws.iter_rows(values_only=True), rows_v, sheets_lc, ext_links, names)
                s["state"] = ws.sheet_state
                try:  # _worksheet_path es privado de openpyxl 3.1; si cambia, el resto del perfil sigue
                    s.update(sheet_layout(path, ws._worksheet_path))
                except Exception as e:  # AttributeError, KeyError o XML ilegible
                    global _layout_warned
                    if not _layout_warned:
                        log.warning("Combinadas/ocultas no disponibles (%s: %s)", type(e).__name__, e)
                        _layout_warned = True
                    s.update(merged_cells=NA, hidden_rows=NA, hidden_columns=NA)
                sheets.append(s)
            except Exception as e:
                log.exception("  falló la hoja '%s'", ws.title)
                failures.append({"sheet": ws.title, "error": f"{type(e).__name__}: {e}"})
    finally:
        wb_f.close()
        if wb_v:
            wb_v.close()
    formulas = sum(s["formulas"] for s in sheets)
    no_cache = sum(s["formulas_without_cache"] for s in sheets)
    cached = "load_failed" if wb_v is None else "n/a" if not formulas else "missing" if formulas and no_cache == formulas else "partial" if no_cache else "ok"
    return {
        "type": path.suffix.lower()[1:], "cached_values": cached, "external_links": ext_links,
        "defined_names": len(names), "sheets": sheets, "sheet_failures": failures,
        "lineage": lineage(sheets, set(sheets_lc.values())),
    }


def csv_rows(path, encoding):
    with open(path, newline="", encoding=encoding) as fh:
        sample = fh.read(65536)
        fh.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        for row in csv.reader(fh, dialect):
            out = []
            for v in row:
                v = v.strip()
                if NUM_RE.match(v):
                    n = v.replace("$", "").replace(",", "")
                    v = float(n) if "." in n else int(n)
                out.append(v)
            yield out


def profile_csv(path):
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            s = analyze_sheet(path.stem, csv_rows(path, enc), None, {}, {}, {})
            break
        except UnicodeDecodeError:
            log.warning("  %s no es %s, reintentando", path.name, enc)
    s.update(state="visible", merged_cells={"total": 0, "ranges": []},
             hidden_rows={"total": 0, "ranges": []}, hidden_columns={"total": 0, "ranges": []})
    return {"type": "csv", "encoding": enc, "cached_values": "n/a", "external_links": {}, "defined_names": 0,
            "sheets": [s], "sheet_failures": [], "lineage": lineage([s], {s["name"]})}


# ---------------------------------------------------------------- reportes

def problems(rep):
    """Lista (prioridad, conteo, texto) ordenada por severidad."""
    out = [(0, 1, f"Hoja **{md(f['sheet'])}** no se pudo procesar: {md(f['error'])}") for f in rep["sheet_failures"]]
    if rep["cached_values"] in ("missing", "load_failed"):
        out.append((1, 1, "El libro no tiene valores cacheados (nunca se recalculó/guardó en Excel): errores y perfiles de columnas calculadas no son confiables"))
    for s in rep["sheets"]:
        n = md(s["name"])
        e = s["errors"]
        if e["total"]:
            by = ", ".join(f"{k}: {v}" for k, v in e["by_type"].items())
            ex = ", ".join(c["cell"] for c in e["cells"][:5])
            out.append((1, e["total"], f"**{n}**: {e['total']} celdas con error ({by}) — p. ej. {ex}"))
        x = s["external_refs"]
        if x["total"]:
            out.append((2, x["total"], f"**{n}**: {x['total']} referencias a libros externos ({md(', '.join(x['books']))})"))
        h = s["hardcoded_formulas"]
        if h["total"]:
            c = h["cells"][0]
            out.append((3, h["total"], f"**{n}**: {h['total']} fórmulas con números fijos — p. ej. {c['cell']} `{md(c['formula'])}`"))
        if s["references"]["unresolved_names"]:
            out.append((4, s["references"]["unresolved_names"], f"**{n}**: {s['references']['unresolved_names']} referencias a nombres/tablas no resueltos"))
        if s["formulas_without_cache"] and rep["cached_values"] == "partial":
            out.append((4, s["formulas_without_cache"], f"**{n}**: {s['formulas_without_cache']} fórmulas sin valor cacheado"))
        if s["state"] != "visible":
            out.append((5, 1, f"**{n}**: hoja con estado `{s['state']}`"))
        if s["merged_cells"] == NA:
            continue
        if s["hidden_rows"]["total"] or s["hidden_columns"]["total"]:
            out.append((5, s["hidden_rows"]["total"] + s["hidden_columns"]["total"],
                        f"**{n}**: {s['hidden_rows']['total']} filas y {s['hidden_columns']['total']} columnas ocultas"))
        if s["merged_cells"]["total"]:
            out.append((6, s["merged_cells"]["total"], f"**{n}**: {s['merged_cells']['total']} rangos combinados"))
    return sorted(out, key=lambda p: (p[0], -p[1]))


def render_md(rep):
    L = [f"# Perfil: {rep['file']}", "",
         f"- Procesado: {rep['processed_at']}",
         f"- Ruta: `{rep['path']}`",
         f"- Tipo: {rep['type']} · Hojas: {len(rep['sheets'])} · Valores cacheados: {rep['cached_values']}"
         f" · Nombres definidos: {rep['defined_names']}", "",
         "## Hojas", "",
         "| Hoja | Estado | Layout | Rango usado | Filas | Cols | No vacías | Fórmulas | % fórm. | Errores | Hardcodes | Externas | Combinadas | Filas ocultas | Cols ocultas | Encabezado |",
         "|---|---|---|---|--:|--:|--:|--:|--:|--:|--:|--:|--:|--:|--:|--:|"]
    tot = lambda f: f["total"] if isinstance(f, dict) else f
    for s in rep["sheets"]:
        L.append(f"| {md(s['name'])} | {s['state']} | {s['layout']} | {s['used_range'] or '—'} | {s['rows']} | {s['columns']} | "
                 f"{s['non_empty_cells']} | {s['formulas']} | {s['formula_pct']} | {s['errors']['total']} | "
                 f"{s['hardcoded_formulas']['total']} | {s['external_refs']['total']} | {tot(s['merged_cells'])} | "
                 f"{tot(s['hidden_rows'])} | {tot(s['hidden_columns'])} | {s['header_row'] or '—'} |")
    for s in rep["sheets"]:
        if s["layout"] == "key_value":
            L += ["", f"### Parámetros: {md(s['name'])} (key_value)", "", "| Fila | Clave | Valor |", "|--:|---|---|"]
            L += [f"| {p['row']} | {md(p['key'] if p['key'] is not None else '')} | "
                  f"{md(p['value'] if p['value'] is not None else '')} |" for p in s["key_values"]]
    L += ["", "## Principales problemas", ""]
    probs = problems(rep)
    L += [f"{i}. {t}" for i, (_, _, t) in enumerate(probs[:15], 1)] or ["Sin problemas detectados."]
    lin = rep["lineage"]
    L += ["", "## Dependencias entre hojas", ""]
    if lin["edges"]:
        ids = {}
        L += ["```mermaid", "flowchart LR"]
        for e in lin["edges"]:
            a, b = (ids.setdefault(x, f"n{len(ids)}") for x in (e["from"], e["to"]))
            label = lambda x: '"' + x.replace('"', "#quot;") + '"'
            L.append(f"  {a}[{label(e['from'])}] -->|{e['refs']}| {b}[{label(e['to'])}]")
        L += ["```", "", "| Origen | Destino | Referencias |", "|---|---|--:|"]
        L += [f"| {md(e['from'])} | {md(e['to'])} | {e['refs']} |" for e in lin["edges"]]
    else:
        L.append("Sin referencias entre hojas.")
    L += ["", f"- **Hojas de entrada** (nadie las alimenta): {', '.join(map(md, lin['input_sheets'])) or '—'}",
          f"- **Hojas de salida** (nadie las consume): {', '.join(map(md, lin['output_sheets'])) or '—'}", "",
          "El detalle por columna (tipos, nulos, únicos, mín/máx, ejemplos) está en el JSON."]
    return "\n".join(L) + "\n"


def readme_section(rep, key):
    sheets = rep["sheets"]
    L = [f"## {rep['file']}", "",
         f"- Fecha: {rep['processed_at']}",
         f"- Hojas: {len(sheets)} · Celdas con contenido: {sum(s['non_empty_cells'] for s in sheets)}"
         f" · Fórmulas: {sum(s['formulas'] for s in sheets)} · Valores cacheados: {rep['cached_values']}",
         "- Hallazgos principales:"]
    L += [f"  - {t}" for _, _, t in problems(rep)[:5]] or ["  - Sin problemas detectados."]
    L.append(f"- Reportes: [JSON]({quote(key + '.profile.json')}) · [Markdown]({quote(key + '.profile.md')})")
    return "\n".join(L)


def update_readme(out, key, section):
    """Anexa o reemplaza solo la sección de este archivo; el resto del README queda intacto."""
    readme = out / "README.md"
    text = readme.read_text(encoding="utf-8") if readme.exists() else \
        "# Reportes de perfilado\n\nUna sección por archivo procesado; se reemplaza al volver a procesarlo.\n"
    start, end = f"<!-- profile:{key} -->", f"<!-- /profile:{key} -->"
    block = f"{start}\n{section}\n{end}"
    pat = re.compile(re.escape(start) + r".*?" + re.escape(end), re.S)
    text = pat.sub(lambda _: block, text, count=1) if pat.search(text) else text.rstrip("\n") + "\n\n" + block + "\n"
    tmp = readme.with_name("README.md.tmp")
    tmp.write_text(text, encoding="utf-8")
    os.replace(tmp, readme)


def process(path, key, out):
    rep = profile_csv(path) if path.suffix.lower() == ".csv" else profile_xlsx(path)
    rep = {"file": key, "path": str(path),
           "processed_at": dt.datetime.now().isoformat(timespec="seconds"), **rep}
    with open(out / f"{key}.profile.json", "w", encoding="utf-8") as fh:
        json.dump(rep, fh, ensure_ascii=False, indent=2, default=str)
    (out / f"{key}.profile.md").write_text(render_md(rep), encoding="utf-8")
    update_readme(out, key, readme_section(rep, key))
    return rep


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("path", nargs="?", default=DATA_DIR, help=f"archivo o carpeta (default: {DATA_DIR})")
    ap.add_argument("--out", default=REPORTS_DIR, help=f"carpeta de reportes (default: {REPORTS_DIR})")
    args = ap.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", datefmt="%H:%M:%S")

    root, out = Path(args.path).expanduser().resolve(), Path(args.out).expanduser().resolve()
    if not root.exists():
        log.error("No existe: %s", root)
        return 2
    if root.is_file():
        files = [root]
    else:
        files = sorted(p for p in root.rglob("*")
                       if p.is_file() and p.suffix.lower() in EXTS and not p.name.startswith("~$"))
    if not files:
        log.warning("No hay archivos .xlsx/.xlsm/.csv en %s", root)
        return 0
    out.mkdir(parents=True, exist_ok=True)

    failed = 0
    for i, f in enumerate(files, 1):
        key = f.name if root.is_file() else "__".join(f.relative_to(root).parts)
        log.info("[%d/%d] %s", i, len(files), f)
        try:
            rep = process(f, key, out)
            log.info("  ok: %d hojas, %d fallas de hoja", len(rep["sheets"]), len(rep["sheet_failures"]))
        except Exception:
            failed += 1
            log.exception("  no se pudo procesar %s", f)
    log.info("Listo: %d archivos, %d con error. Reportes en %s", len(files), failed, out)
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
